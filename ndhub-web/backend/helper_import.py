"""Domain helpers for ND-Hub web backend (Issue #96)."""

from __future__ import annotations

import hashlib
import json
import logging
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

from fastapi import UploadFile

try:
    import pandas as pd
except ImportError:  # pragma: no cover
    pd = None
try:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    Workbook = None
    DataValidation = None

from backend.database import SqliteRepository
from backend.helper_common import _sanitize_filename_part
from backend.models import (
    ALLOWED_IMPORT_TYPES,
    IMPORT_COLUMN_ALIASES,
    IMPORT_REQUIRED_COLUMNS,
    MAX_ATTACHMENT_SIZE_BYTES,
)


def _normalize_import_column_name(name: str) -> str:
    lowered = (name or "").strip().lower()
    lowered = lowered.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    lowered = re.sub(r"[^a-z0-9]+", "", lowered)
    return lowered

def _ensure_import_dependencies() -> None:
    if pd is None:
        raise ValueError("Import-Feature benoetigt pandas. Bitte Abhaengigkeit installieren.")
    if Workbook is None or DataValidation is None:
        raise ValueError("Import-Feature benoetigt openpyxl. Bitte Abhaengigkeit installieren.")

def _load_import_dataframe(upload_file: UploadFile) -> tuple[pd.DataFrame, int]:
    _ensure_import_dependencies()
    filename = (upload_file.filename or "").lower()
    upload_file.file.seek(0)
    if filename.endswith(".csv"):
        df = pd.read_csv(upload_file.file)
        return df, 2
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        upload_file.file.seek(0)
        df = pd.read_excel(upload_file.file, header=2)
        normalized_columns = {_normalize_import_column_name(str(col)) for col in df.columns}
        normalized_required = {
            _normalize_import_column_name(alias)
            for alias in IMPORT_COLUMN_ALIASES["depot"]
        }
        if not normalized_required.intersection(normalized_columns):
            upload_file.file.seek(0)
            df = pd.read_excel(upload_file.file, header=0)
            return df, 2
        return df, 4
    raise ValueError("Nur CSV-, XLSX- oder XLS-Dateien sind erlaubt.")

def _map_import_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename_map: dict[str, str] = {}
    for source_col in df.columns:
        normalized = _normalize_import_column_name(str(source_col))
        for target_col, aliases in IMPORT_COLUMN_ALIASES.items():
            normalized_aliases = {_normalize_import_column_name(alias) for alias in aliases}
            if normalized in normalized_aliases:
                rename_map[source_col] = target_col
                break
    mapped = df.rename(columns=rename_map).copy()
    missing = [col for col in IMPORT_REQUIRED_COLUMNS if col not in mapped.columns]
    if missing:
        raise ValueError(f"Fehlende Spalten: {', '.join(missing)}")
    return mapped

def _analyze_import_dataframe(
    repository: SqliteRepository,
    mapped_df: pd.DataFrame,
    row_offset: int,
) -> tuple[list[dict], list[dict], list[str]]:
    working = mapped_df.dropna(how="all").copy()
    if "depot" in working.columns:
        working = working[working["depot"].notna()]
        working = working[working["depot"].astype(str).str.strip() != ""]

    depots = repository.list_depots(limit=10000, offset=0, q="")
    depots_by_name = {str(row["name"]).strip().lower(): int(row["id"]) for row in depots}
    praeparate = repository.list_praeparate(limit=10000, offset=0, q="")
    prae_by_name = {str(row["name"]).strip().lower(): int(row["id"]) for row in praeparate}

    parsed_rows: list[dict] = []
    errors: list[str] = []

    for idx, row in working.iterrows():
        display_row = int(idx) + row_offset
        depot_name = str(row.get("depot", "")).strip()
        prae_name = str(row.get("praeparat", "")).strip()
        typ = str(row.get("typ", "")).strip()
        charge = str(row.get("charge", "")).strip()
        empfaenger_raw = row.get("empfaenger", None)
        empfaenger = str(empfaenger_raw).strip() if pd.notna(empfaenger_raw) else None

        depot_id = depots_by_name.get(depot_name.lower())
        if depot_id is None:
            errors.append(f"Zeile {display_row}: Depot '{depot_name}' nicht gefunden")
            continue
        praeparat_id = prae_by_name.get(prae_name.lower())
        if praeparat_id is None:
            errors.append(f"Zeile {display_row}: Praeparat '{prae_name}' nicht gefunden")
            continue
        if typ not in ALLOWED_IMPORT_TYPES:
            errors.append(f"Zeile {display_row}: Ungueltiger Typ '{typ}'")
            continue
        if not charge:
            errors.append(f"Zeile {display_row}: Charge darf nicht leer sein")
            continue
        try:
            anzahl = int(row.get("anzahl"))
            if anzahl <= 0:
                raise ValueError()
        except (ValueError, TypeError):
            errors.append(f"Zeile {display_row}: Anzahl muss > 0 sein")
            continue
        datum_iso = _parse_import_date(row.get("datum"))
        if datum_iso is None:
            errors.append(f"Zeile {display_row}: Datum ist ungueltig")
            continue
        verfall_iso = _parse_import_date(row.get("verfall"))
        if verfall_iso is None:
            errors.append(f"Zeile {display_row}: Verfall ist ungueltig")
            continue
        parsed_rows.append(
            {
                "display_row": display_row,
                "depot_id": depot_id,
                "praeparat_id": praeparat_id,
                "typ": typ,
                "charge": charge,
                "verfall": verfall_iso,
                "datum": datum_iso,
                "anzahl": anzahl,
                "empfaenger": empfaenger or None,
                "depot": depot_name,
                "praeparat": prae_name,
            }
        )

    preview_rows = working.head(10).fillna("").to_dict(orient="records")
    for row in preview_rows:
        for key, value in row.items():
            row[key] = str(value)
    return parsed_rows, preview_rows, errors

def _parse_import_date(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    value_text = str(value).strip()
    if not value_text:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value_text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    parsed = pd.to_datetime(value_text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.strftime("%Y-%m-%d")

def _build_import_execution_fingerprint(filename: str, parsed_rows: list[dict[str, Any]]) -> str:
    canonical_rows: list[dict[str, Any]] = []
    for row in parsed_rows:
        canonical_rows.append(
            {
                "depot_id": int(row.get("depot_id") or 0),
                "praeparat_id": int(row.get("praeparat_id") or 0),
                "typ": str(row.get("typ") or ""),
                "charge": str(row.get("charge") or ""),
                "verfall": str(row.get("verfall") or ""),
                "datum": str(row.get("datum") or ""),
                "anzahl": int(row.get("anzahl") or 0),
                "empfaenger": str(row.get("empfaenger") or ""),
            }
        )
    payload = {
        "filename": str(filename or "").strip().lower(),
        "rows": canonical_rows,
    }
    raw = json.dumps(payload, ensure_ascii=True, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()

def _classify_import_error_code(error_message: str) -> str:
    text = str(error_message or "").lower()
    if "depot" in text and "nicht gefunden" in text:
        return "unknown_depot"
    if "praeparat" in text and "nicht gefunden" in text:
        return "unknown_praeparat"
    if "ungueltiger typ" in text:
        return "invalid_type"
    if "charge" in text and "leer" in text:
        return "missing_charge"
    if "anzahl" in text:
        return "invalid_amount"
    if "datum" in text:
        return "invalid_date"
    if "verfall" in text:
        return "invalid_expiry"
    if "duplikat" in text:
        return "duplicate"
    return "other"

def _summarize_import_errors(errors: list[str]) -> dict[str, Any]:
    by_code: dict[str, int] = {}
    for item in errors:
        code = _classify_import_error_code(item)
        by_code[code] = by_code.get(code, 0) + 1
    return {"total": int(len(errors)), "by_code": by_code}

def _build_import_template(repository: SqliteRepository) -> bytes:
    _ensure_import_dependencies()
    wb = Workbook()
    ws = wb.active
    ws.title = "Bewegungen"
    depots = [row["name"] for row in repository.list_depots(limit=10000, offset=0, q="")]
    praeparate = [row["name"] for row in repository.list_praeparate(limit=10000, offset=0, q="")]
    typen = ["Zugang", "Abgang", "Vernichtung"]

    validation = wb.create_sheet("Validierung")
    for idx, name in enumerate(depots, start=1):
        validation.cell(row=idx, column=1, value=name)
    for idx, name in enumerate(praeparate, start=1):
        validation.cell(row=idx, column=2, value=name)
    for idx, name in enumerate(typen, start=1):
        validation.cell(row=idx, column=3, value=name)
    validation.sheet_state = "hidden"

    ws.merge_cells("A1:H1")
    ws["A1"] = "Depot in B2 waehlen, danach Daten ab Zeile 4 eintragen."
    ws["A2"] = "Depot:"
    ws["B2"] = depots[0] if depots else ""
    headers = ["Depot", "Praeparat", "Typ", "Charge", "Verfall", "Datum", "Anzahl", "Empfaenger"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=header)
    for row_idx in range(4, 104):
        ws.cell(row=row_idx, column=1, value=f'=IF(B{row_idx}<>"",$B$2,"")')

    if depots:
        dv_depot = DataValidation(type="list", formula1=f"Validierung!$A$1:$A${len(depots)}", allow_blank=False)
        ws.add_data_validation(dv_depot)
        dv_depot.add("B2")
    if praeparate:
        dv_praep = DataValidation(type="list", formula1=f"Validierung!$B$1:$B${len(praeparate)}", allow_blank=False)
        ws.add_data_validation(dv_praep)
        dv_praep.add("B4:B103")
    dv_typ = DataValidation(type="list", formula1=f"Validierung!$C$1:$C${len(typen)}", allow_blank=False)
    ws.add_data_validation(dv_typ)
    dv_typ.add("C4:C103")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def _build_attachment_target_path(
    base_dir: Path,
    bewegung_id: int,
    depot_name: str | None,
    original_name: str,
) -> Path:
    now = datetime.now()
    folder = base_dir / _sanitize_filename_part(depot_name or "depot") / now.strftime("%Y") / now.strftime("%m")
    folder.mkdir(parents=True, exist_ok=True)
    timestamp = now.strftime("%Y%m%d_%H%M%S")
    stem = _sanitize_filename_part(Path(original_name).stem or "anhang")
    filename = f"bewegung_{bewegung_id}_{timestamp}_{stem}.pdf"
    return folder / filename

def _persist_pdf_upload(upload_file: UploadFile, target: Path) -> int:
    source = upload_file.file
    source.seek(0)
    total_written = 0
    with target.open("wb") as out:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            total_written += len(chunk)
            if total_written > MAX_ATTACHMENT_SIZE_BYTES:
                out.close()
                target.unlink(missing_ok=True)
                raise ValueError("PDF-Datei ist zu gross (max. 10 MB).")
            out.write(chunk)
    return total_written

