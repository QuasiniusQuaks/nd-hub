"""Domain helpers for ND-Hub web backend (Issue #96)."""

from __future__ import annotations

import csv
import io
import logging
import tempfile
from datetime import UTC, datetime
from io import BytesIO

logger = logging.getLogger(__name__)

from fastapi import HTTPException, status
from fastapi.responses import StreamingResponse

try:
    import pandas as pd
except ImportError:  # pragma: no cover
    pd = None
try:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    Workbook = None
    DataValidation = None


from backend.helper_common import (
    _parse_optional_iso_date,
    _safe_report_filename_token,
)


def _normalize_verfall_thresholds(
    critical_days: int = 30,
    warning_days: int = 90,
    attention_days: int = 180,
) -> tuple[int, int, int]:
    critical = max(1, min(int(critical_days), 3650))
    warning = max(critical + 1, min(int(warning_days), 3650))
    attention = max(warning + 1, min(int(attention_days), 3650))
    return critical, warning, attention

def _verfall_category(
    tage_bis_verfall: int,
    critical_days: int,
    warning_days: int,
    attention_days: int,
) -> str:
    if tage_bis_verfall <= critical_days:
        return "kritisch"
    if tage_bis_verfall <= warning_days:
        return "warnung"
    if tage_bis_verfall <= attention_days:
        return "achtung"
    return "ok"

def _enrich_verfall_rows(
    rows: list[dict],
    critical_days: int,
    warning_days: int,
    attention_days: int,
) -> list[dict]:
    enriched: list[dict] = []
    for row in rows:
        item = dict(row)
        try:
            tage = int(item.get("tage_bis_verfall"))
        except (TypeError, ValueError):
            tage = 99999
        item["tage_bis_verfall"] = tage
        item["kategorie"] = _verfall_category(tage, critical_days, warning_days, attention_days)
        enriched.append(item)
    return enriched

def _validated_report_date_range(start_date: str, end_date: str) -> tuple[str | None, str | None]:
    safe_start = _parse_optional_iso_date(start_date, "start_date")
    safe_end = _parse_optional_iso_date(end_date, "end_date")
    if safe_start and safe_end and safe_start > safe_end:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="start_date darf nicht nach end_date liegen.")
    return safe_start, safe_end

def _build_report_export_filename(
    report_slug: str,
    extension: str,
    perspective: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    horizon_months: int | None = None,
    limit: int | None = None,
) -> str:
    parts = [_safe_report_filename_token(report_slug)]
    safe_perspective = _safe_report_filename_token(perspective or "")
    if safe_perspective and safe_perspective in {"depot", "praeparat"}:
        parts.append(safe_perspective)
    if start_date or end_date:
        start_token = _safe_report_filename_token(start_date or "all")
        end_token = _safe_report_filename_token(end_date or "all")
        parts.append(f"{start_token}_to_{end_token}")
    if horizon_months is not None:
        parts.append(f"h{max(1, int(horizon_months))}m")
    if limit is not None:
        parts.append(f"top{max(1, int(limit))}")
    stamp = datetime.now(UTC).strftime("%Y%m%d")
    safe_extension = _safe_report_filename_token(extension) or "csv"
    return f"{'_'.join(parts)}_{stamp}.{safe_extension}"

def _csv_response(filename: str, headers: list[str], rows: list[list[object]]) -> StreamingResponse:
    text_io = io.StringIO()
    writer = csv.writer(text_io)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    buffer = BytesIO(text_io.getvalue().encode("utf-8"))
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

def _pdf_response(filename: str, title: str, headers: list[str], rows: list[list[object]], kpis: dict) -> StreamingResponse:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PDF-Export erfordert reportlab.") from exc

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Heading1"]), Spacer(1, 8)]
    if kpis:
        for key, value in kpis.items():
            elements.append(Paragraph(f"<b>{key}:</b> {value}", styles["Normal"]))
        elements.append(Spacer(1, 8))
    table_data = [headers] + rows
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#007aff")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

def _pptx_response(filename: str, title: str, headers: list[str], rows: list[list[object]], kpis: dict) -> StreamingResponse:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PPT-Export erfordert python-pptx.") from exc

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_box = title_slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.size = Pt(24)
    title_tf.paragraphs[0].font.bold = True

    y = 1.0
    if kpis:
        kpi_box = title_slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(8.8), Inches(1.4))
        kpi_tf = kpi_box.text_frame
        first = True
        for key, value in kpis.items():
            p = kpi_tf.paragraphs[0] if first else kpi_tf.add_paragraph()
            first = False
            p.text = f"{key}: {value}"
            p.font.size = Pt(14)
        y += 1.5

    max_rows = min(len(rows), 18)
    cols = len(headers)
    table_shape = title_slide.shapes.add_table(cols=cols, rows=max_rows + 1, left=Inches(0.4), top=Inches(y), width=Inches(9.0), height=Inches(5.0))
    table = table_shape.table
    for col_idx, header in enumerate(headers):
        table.cell(0, col_idx).text = str(header)
    for r in range(max_rows):
        for c in range(cols):
            table.cell(r + 1, c).text = str(rows[r][c])

    with tempfile.NamedTemporaryFile(suffix=".pptx") as tmp:
        prs.save(tmp.name)
        tmp.seek(0)
        payload = tmp.read()
    out = BytesIO(payload)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

