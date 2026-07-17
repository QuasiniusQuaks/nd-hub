"""FastAPI app for ND-Hub MVP backend."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import re
import shutil
import smtplib
import sqlite3
import tempfile
from contextlib import asynccontextmanager
from datetime import UTC, date, datetime
from decimal import Decimal
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path
from typing import Any, Optional
from urllib.parse import urlencode, urlparse
from urllib.request import Request as UrlRequest
from urllib.request import urlopen

logger = logging.getLogger(__name__)

from fastapi import (
    Depends,
    FastAPI,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
    status,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

# ---- Rate-Limit fuer /auth/login (Fix fuer #32) ----
from slowapi import Limiter
from slowapi.errors import RateLimitExceeded
from slowapi.util import get_remote_address
from starlette.requests import Request as _StarletteRequest
from starlette.responses import JSONResponse as _JSONResponse

# ---- /Rate-Limit ----
try:
    import pandas as pd
except ImportError:  # pragma: no cover - optional runtime dependency
    pd = None
try:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - optional runtime dependency
    Workbook = None
    DataValidation = None

from backend.auth import (
    SessionInfo,
    TokenStore,
    bearer_scheme,
    get_current_session,
)
from backend.config import (
    EmailDeliverySettings,
    resolve_auto_backup_hours,
    resolve_db_engine,
    resolve_dual_write_sqlite,
    resolve_email_delivery_settings,
    resolve_mariadb_settings,
    resolve_max_backup_restore_mb,
    resolve_runtime_paths,
)
from backend.database import SqliteRepository
from backend.mariadb_repository import MariaDbRepository
from security_manager import SecurityManager


def _require_http_scheme(url: str) -> str:
    """Validates that the URL uses only http/https (SSRF protection)."""
    scheme = urlparse(url).scheme.lower()
    if scheme not in {"http", "https"}:
        raise ValueError(f"URL scheme not allowed: {url}")
    return url


class LoginRequest(BaseModel):
    username: str = Field(min_length=1)
    password: str = Field(min_length=1)


class LoginResponse(BaseModel):
    token: str
    username: str
    role: Optional[str]
    requires_password_change: bool = False
    permissions: list[str] = Field(default_factory=list)


class DesktopSyncTokenResponse(BaseModel):
    backend_url: str
    token: str
    username: str
    role: Optional[str]
    expires_at: str
    client_label: Optional[str] = None


class DesktopSyncTokenCreateRequest(BaseModel):
    client_label: Optional[str] = Field(default=None, max_length=120)


class DesktopSyncTokenArchiveItem(BaseModel):
    token_fingerprint: str
    token_masked: str
    client_label: Optional[str]
    username: str
    role: Optional[str]
    created_at: str
    expires_at: str
    status: str


class DesktopSyncTokenRevokeRequest(BaseModel):
    token_fingerprint: str = Field(min_length=8)


class PasswordChangeRequest(BaseModel):
    old_password: str = Field(min_length=1)
    new_password: str = Field(min_length=8)


class UserCreateRequest(BaseModel):
    username: str = Field(min_length=3)
    password: str = Field(min_length=8)
    role: str = Field(min_length=1)
    email: Optional[str] = None
    is_active: bool = True
    permissions: Optional[list[str]] = None


class UserUpdateRequest(BaseModel):
    username: Optional[str] = Field(default=None, min_length=3)
    role: Optional[str] = None
    email: Optional[str] = None
    is_active: Optional[bool] = None
    permissions: Optional[list[str]] = None


class UserPasswordResetRequest(BaseModel):
    new_password: str = Field(min_length=8)


class BewegungCreateRequest(BaseModel):
    depot_id: int
    praeparat_id: int
    typ: str
    charge: str
    verfall: date
    datum: date
    anzahl: int = Field(gt=0)
    empfaenger: Optional[str] = None


class DepotUpsertRequest(BaseModel):
    name: str = Field(min_length=1)
    adresse: Optional[str] = None
    strasse: Optional[str] = None
    hausnummer: Optional[str] = None
    postleitzahl: Optional[str] = None
    stadt: Optional[str] = None
    telefon: Optional[str] = None
    email: Optional[str] = None
    institution_id: Optional[int] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None


class InstitutionUpsertRequest(BaseModel):
    name: str = Field(min_length=1)
    adresse: Optional[str] = None
    strasse: Optional[str] = None
    hausnummer: Optional[str] = None
    postleitzahl: Optional[str] = None
    stadt: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None


class OnboardingPraeparatItem(BaseModel):
    name: str = Field(min_length=1)
    wirkstoff: Optional[str] = None
    darreichungsform: Optional[str] = None
    staerke: Optional[str] = None
    einheit: Optional[str] = None
    pzn: Optional[str] = None
    hersteller: Optional[str] = None


class OnboardingDepotAssignmentItem(BaseModel):
    praeparat_name: str = Field(min_length=1)
    sollbestand: int = Field(ge=0, default=0)


class OnboardingDepotItem(BaseModel):
    name: str = Field(min_length=1)
    adresse: Optional[str] = None
    strasse: Optional[str] = None
    hausnummer: Optional[str] = None
    postleitzahl: Optional[str] = None
    stadt: Optional[str] = None
    telefon: Optional[str] = None
    email: Optional[str] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    assignments: list[OnboardingDepotAssignmentItem] = Field(default_factory=list)


class OnboardingInstitutionSetupRequest(BaseModel):
    institution: InstitutionUpsertRequest
    praeparate: list[OnboardingPraeparatItem] = Field(default_factory=list)
    depots: list[OnboardingDepotItem] = Field(default_factory=list)


class UserDepotPermissionItem(BaseModel):
    depot_id: int
    can_read: bool = False
    can_write: bool = False


class UserDepotPermissionUpdateRequest(BaseModel):
    username: str = Field(min_length=1)
    permissions: list[UserDepotPermissionItem] = Field(default_factory=list)


class PraeparatUpsertRequest(BaseModel):
    name: str = Field(min_length=1)
    wirkstoff: Optional[str] = None
    darreichungsform: Optional[str] = None
    staerke: Optional[str] = None
    einheit: Optional[str] = None
    pzn: Optional[str] = None
    hersteller: Optional[str] = None


class DepotAssignmentItem(BaseModel):
    praeparat_id: int
    sollbestand: int = Field(ge=0, default=0)


class DepotAssignmentsUpdateRequest(BaseModel):
    assignments: list[DepotAssignmentItem]


class KontaktUpsertRequest(BaseModel):
    name: str = Field(min_length=1)
    rolle: Optional[str] = None
    telefon: Optional[str] = None
    email: Optional[str] = None


class EmailRecipientPreviewRequest(BaseModel):
    depot_ids: list[int] = Field(default_factory=list)
    kontakt_ids: list[int] = Field(default_factory=list)


class EmailDraftCreateRequest(BaseModel):
    depot_ids: list[int] = Field(default_factory=list)
    kontakt_ids: list[int] = Field(default_factory=list)
    betreff: str = Field(min_length=1)
    nachricht: str = ""
    send_now: bool = False


class EmailDeliveryStatusUpdateRequest(BaseModel):
    versand_status: str = Field(min_length=1)
    versand_kanal: Optional[str] = None
    versand_fehler: Optional[str] = None


class SyncPullRequest(BaseModel):
    cursor: Optional[str] = None
    entities: list[str] = Field(default_factory=list)
    limit: int = Field(default=200, ge=1, le=1000)


class SyncPushChangeItem(BaseModel):
    entity: str = Field(min_length=1)
    operation: str = Field(min_length=1)
    payload: dict[str, Any] = Field(default_factory=dict)
    client_change_id: Optional[str] = None


class SyncPushRequest(BaseModel):
    batch_id: str = Field(min_length=1, max_length=128)
    changes: list[SyncPushChangeItem] = Field(default_factory=list)


MAX_ATTACHMENT_SIZE_BYTES = 10 * 1024 * 1024
PDF_MIME_TYPES = {"application/pdf", "application/x-pdf"}
ALLOWED_IMPORT_TYPES = {"Zugang", "Abgang", "Vernichtung"}
IMPORT_REQUIRED_COLUMNS = ["depot", "praeparat", "typ", "charge", "verfall", "datum", "anzahl"]
IMPORT_COLUMN_ALIASES = {
    "depot": {"depot"},
    "praeparat": {"praeparat", "präparat"},
    "typ": {"typ", "type"},
    "charge": {"charge"},
    "verfall": {"verfall", "verfallsdatum"},
    "datum": {"datum", "date"},
    "anzahl": {"anzahl", "menge", "quantity"},
    "empfaenger": {"empfaenger", "empfänger"},
}


def _email_delivery_missing_config(settings: EmailDeliverySettings) -> list[str]:
    if settings.mode != "smtp":
        return []
    missing: list[str] = []
    if not settings.smtp_host:
        missing.append("ND_HUB_SMTP_HOST")
    if not settings.smtp_from_address:
        missing.append("ND_HUB_SMTP_FROM_ADDRESS")
    if settings.smtp_use_ssl and settings.smtp_use_tls:
        missing.append("ND_HUB_SMTP_USE_TLS/ND_HUB_SMTP_USE_SSL")
    return missing


def _send_email_via_smtp(
    settings: EmailDeliverySettings,
    subject: str,
    message: str,
    recipients: list[str],
) -> dict[str, Any]:
    if settings.mode != "smtp":
        raise RuntimeError("Live-Versand ist deaktiviert (ND_HUB_EMAIL_DELIVERY_MODE=draft).")
    missing = _email_delivery_missing_config(settings)
    if missing:
        joined = ", ".join(missing)
        raise RuntimeError(f"SMTP-Konfiguration unvollstaendig: {joined}")
    if not recipients:
        raise RuntimeError("Keine Empfaenger vorhanden.")

    msg = EmailMessage()
    msg["Subject"] = (subject or "").strip() or "ND-Hub Nachricht"
    msg["From"] = (
        f"{settings.smtp_from_name} <{settings.smtp_from_address}>"
        if settings.smtp_from_name
        else settings.smtp_from_address
    )
    msg["To"] = ", ".join(recipients)
    msg.set_content(message or "")

    smtp_timeout = max(3, int(settings.smtp_timeout_seconds))
    try:
        if settings.smtp_use_ssl:
            smtp_client = smtplib.SMTP_SSL(settings.smtp_host, settings.smtp_port, timeout=smtp_timeout)
        else:
            smtp_client = smtplib.SMTP(settings.smtp_host, settings.smtp_port, timeout=smtp_timeout)
        with smtp_client as smtp:
            smtp.ehlo()
            if settings.smtp_use_tls:
                smtp.starttls()
                smtp.ehlo()
            if settings.smtp_username:
                smtp.login(settings.smtp_username, settings.smtp_password)
            rejected = smtp.send_message(msg) or {}
    except Exception as exc:
        raise RuntimeError(f"SMTP-Versand fehlgeschlagen: {exc}") from exc

    rejected_addresses = sorted(str(address) for address in rejected.keys())
    sent_count = max(0, len(recipients) - len(rejected_addresses))
    if sent_count <= 0:
        raise RuntimeError("SMTP hat keine Empfaenger akzeptiert.")
    return {"sent_count": sent_count, "rejected_recipients": rejected_addresses}


def _sanitize_filename_part(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", value or "")
    return safe.strip("._") or "datei"


def _quote_sql_identifier(name: str) -> str:
    """Quote and validate a SQL identifier (table/column name)."""
    if not name or not re.fullmatch(r"[A-Za-z0-9_]+", name):
        raise ValueError(f"Ungueltiger SQL-Bezeichner: {name!r}")
    return f"`{name}`"


def _build_attachment_target_path(
    base_dir: Path,
    bewegung_id: int,
    depot_name: str | None,
    original_name: str,
) -> Path:
    now = datetime.now()
    folder = base_dir / _sanitize_filename_part(depot_name or "depot") / now.strftime("%Y") / now.strftime("%m")
    folder.mkdir(parents=True, exist_ok=True)
    timestamp = now.strftime("%Y%m%d_%H%M%S")
    stem = _sanitize_filename_part(Path(original_name).stem or "anhang")
    filename = f"bewegung_{bewegung_id}_{timestamp}_{stem}.pdf"
    return folder / filename


def _persist_pdf_upload(upload_file: UploadFile, target: Path) -> int:
    source = upload_file.file
    source.seek(0)
    total_written = 0
    with target.open("wb") as out:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            total_written += len(chunk)
            if total_written > MAX_ATTACHMENT_SIZE_BYTES:
                out.close()
                target.unlink(missing_ok=True)
                raise ValueError("PDF-Datei ist zu gross (max. 10 MB).")
            out.write(chunk)
    return total_written


def _normalize_import_column_name(name: str) -> str:
    lowered = (name or "").strip().lower()
    lowered = lowered.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    lowered = re.sub(r"[^a-z0-9]+", "", lowered)
    return lowered


def _parse_id_list_csv(value: str) -> list[int]:
    if not value.strip():
        return []
    result: list[int] = []
    for part in value.split(","):
        item = part.strip()
        if not item:
            continue
        result.append(int(item))
    return result


def _add_months(iso_date: date, months: int) -> date:
    month_index = (iso_date.year * 12 + iso_date.month - 1) + months
    year = month_index // 12
    month = (month_index % 12) + 1
    return date(year, month, 1)


def _normalize_verfall_thresholds(
    critical_days: int = 30,
    warning_days: int = 90,
    attention_days: int = 180,
) -> tuple[int, int, int]:
    critical = max(1, min(int(critical_days), 3650))
    warning = max(critical + 1, min(int(warning_days), 3650))
    attention = max(warning + 1, min(int(attention_days), 3650))
    return critical, warning, attention


def _verfall_category(
    tage_bis_verfall: int,
    critical_days: int,
    warning_days: int,
    attention_days: int,
) -> str:
    if tage_bis_verfall <= critical_days:
        return "kritisch"
    if tage_bis_verfall <= warning_days:
        return "warnung"
    if tage_bis_verfall <= attention_days:
        return "achtung"
    return "ok"


def _enrich_verfall_rows(
    rows: list[dict],
    critical_days: int,
    warning_days: int,
    attention_days: int,
) -> list[dict]:
    enriched: list[dict] = []
    for row in rows:
        item = dict(row)
        try:
            tage = int(item.get("tage_bis_verfall"))
        except (TypeError, ValueError):
            tage = 99999
        item["tage_bis_verfall"] = tage
        item["kategorie"] = _verfall_category(tage, critical_days, warning_days, attention_days)
        enriched.append(item)
    return enriched


def _ensure_import_dependencies() -> None:
    if pd is None:
        raise ValueError("Import-Feature benoetigt pandas. Bitte Abhaengigkeit installieren.")
    if Workbook is None or DataValidation is None:
        raise ValueError("Import-Feature benoetigt openpyxl. Bitte Abhaengigkeit installieren.")


def _load_import_dataframe(upload_file: UploadFile) -> tuple[pd.DataFrame, int]:
    _ensure_import_dependencies()
    filename = (upload_file.filename or "").lower()
    upload_file.file.seek(0)
    if filename.endswith(".csv"):
        df = pd.read_csv(upload_file.file)
        return df, 2
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        upload_file.file.seek(0)
        df = pd.read_excel(upload_file.file, header=2)
        normalized_columns = {_normalize_import_column_name(str(col)) for col in df.columns}
        normalized_required = {
            _normalize_import_column_name(alias)
            for alias in IMPORT_COLUMN_ALIASES["depot"]
        }
        if not normalized_required.intersection(normalized_columns):
            upload_file.file.seek(0)
            df = pd.read_excel(upload_file.file, header=0)
            return df, 2
        return df, 4
    raise ValueError("Nur CSV-, XLSX- oder XLS-Dateien sind erlaubt.")


def _map_import_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename_map: dict[str, str] = {}
    for source_col in df.columns:
        normalized = _normalize_import_column_name(str(source_col))
        for target_col, aliases in IMPORT_COLUMN_ALIASES.items():
            normalized_aliases = {_normalize_import_column_name(alias) for alias in aliases}
            if normalized in normalized_aliases:
                rename_map[source_col] = target_col
                break
    mapped = df.rename(columns=rename_map).copy()
    missing = [col for col in IMPORT_REQUIRED_COLUMNS if col not in mapped.columns]
    if missing:
        raise ValueError(f"Fehlende Spalten: {', '.join(missing)}")
    return mapped


def _analyze_import_dataframe(
    repository: SqliteRepository,
    mapped_df: pd.DataFrame,
    row_offset: int,
) -> tuple[list[dict], list[dict], list[str]]:
    working = mapped_df.dropna(how="all").copy()
    if "depot" in working.columns:
        working = working[working["depot"].notna()]
        working = working[working["depot"].astype(str).str.strip() != ""]

    depots = repository.list_depots(limit=10000, offset=0, q="")
    depots_by_name = {str(row["name"]).strip().lower(): int(row["id"]) for row in depots}
    praeparate = repository.list_praeparate(limit=10000, offset=0, q="")
    prae_by_name = {str(row["name"]).strip().lower(): int(row["id"]) for row in praeparate}

    parsed_rows: list[dict] = []
    errors: list[str] = []

    for idx, row in working.iterrows():
        display_row = int(idx) + row_offset
        depot_name = str(row.get("depot", "")).strip()
        prae_name = str(row.get("praeparat", "")).strip()
        typ = str(row.get("typ", "")).strip()
        charge = str(row.get("charge", "")).strip()
        empfaenger_raw = row.get("empfaenger", None)
        empfaenger = str(empfaenger_raw).strip() if pd.notna(empfaenger_raw) else None

        depot_id = depots_by_name.get(depot_name.lower())
        if depot_id is None:
            errors.append(f"Zeile {display_row}: Depot '{depot_name}' nicht gefunden")
            continue
        praeparat_id = prae_by_name.get(prae_name.lower())
        if praeparat_id is None:
            errors.append(f"Zeile {display_row}: Praeparat '{prae_name}' nicht gefunden")
            continue
        if typ not in ALLOWED_IMPORT_TYPES:
            errors.append(f"Zeile {display_row}: Ungueltiger Typ '{typ}'")
            continue
        if not charge:
            errors.append(f"Zeile {display_row}: Charge darf nicht leer sein")
            continue
        try:
            anzahl = int(row.get("anzahl"))
            if anzahl <= 0:
                raise ValueError()
        except (ValueError, TypeError):
            errors.append(f"Zeile {display_row}: Anzahl muss > 0 sein")
            continue
        datum_iso = _parse_import_date(row.get("datum"))
        if datum_iso is None:
            errors.append(f"Zeile {display_row}: Datum ist ungueltig")
            continue
        verfall_iso = _parse_import_date(row.get("verfall"))
        if verfall_iso is None:
            errors.append(f"Zeile {display_row}: Verfall ist ungueltig")
            continue
        parsed_rows.append(
            {
                "display_row": display_row,
                "depot_id": depot_id,
                "praeparat_id": praeparat_id,
                "typ": typ,
                "charge": charge,
                "verfall": verfall_iso,
                "datum": datum_iso,
                "anzahl": anzahl,
                "empfaenger": empfaenger or None,
                "depot": depot_name,
                "praeparat": prae_name,
            }
        )

    preview_rows = working.head(10).fillna("").to_dict(orient="records")
    for row in preview_rows:
        for key, value in row.items():
            row[key] = str(value)
    return parsed_rows, preview_rows, errors


def _parse_import_date(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    value_text = str(value).strip()
    if not value_text:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value_text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    parsed = pd.to_datetime(value_text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.strftime("%Y-%m-%d")


def _parse_optional_iso_date(value: str, field_name: str) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"{field_name} muss YYYY-MM-DD sein.") from exc


def _validated_report_date_range(start_date: str, end_date: str) -> tuple[str | None, str | None]:
    safe_start = _parse_optional_iso_date(start_date, "start_date")
    safe_end = _parse_optional_iso_date(end_date, "end_date")
    if safe_start and safe_end and safe_start > safe_end:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="start_date darf nicht nach end_date liegen.")
    return safe_start, safe_end


def _safe_report_filename_token(value: str) -> str:
    return re.sub(r"[^a-z0-9_-]+", "_", (value or "").strip().lower()).strip("_") or "report"


def _build_report_export_filename(
    report_slug: str,
    extension: str,
    perspective: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    horizon_months: int | None = None,
    limit: int | None = None,
) -> str:
    parts = [_safe_report_filename_token(report_slug)]
    safe_perspective = _safe_report_filename_token(perspective or "")
    if safe_perspective and safe_perspective in {"depot", "praeparat"}:
        parts.append(safe_perspective)
    if start_date or end_date:
        start_token = _safe_report_filename_token(start_date or "all")
        end_token = _safe_report_filename_token(end_date or "all")
        parts.append(f"{start_token}_to_{end_token}")
    if horizon_months is not None:
        parts.append(f"h{max(1, int(horizon_months))}m")
    if limit is not None:
        parts.append(f"top{max(1, int(limit))}")
    stamp = datetime.now(UTC).strftime("%Y%m%d")
    safe_extension = _safe_report_filename_token(extension) or "csv"
    return f"{'_'.join(parts)}_{stamp}.{safe_extension}"


def _build_import_execution_fingerprint(filename: str, parsed_rows: list[dict[str, Any]]) -> str:
    canonical_rows: list[dict[str, Any]] = []
    for row in parsed_rows:
        canonical_rows.append(
            {
                "depot_id": int(row.get("depot_id") or 0),
                "praeparat_id": int(row.get("praeparat_id") or 0),
                "typ": str(row.get("typ") or ""),
                "charge": str(row.get("charge") or ""),
                "verfall": str(row.get("verfall") or ""),
                "datum": str(row.get("datum") or ""),
                "anzahl": int(row.get("anzahl") or 0),
                "empfaenger": str(row.get("empfaenger") or ""),
            }
        )
    payload = {
        "filename": str(filename or "").strip().lower(),
        "rows": canonical_rows,
    }
    raw = json.dumps(payload, ensure_ascii=True, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _classify_import_error_code(error_message: str) -> str:
    text = str(error_message or "").lower()
    if "depot" in text and "nicht gefunden" in text:
        return "unknown_depot"
    if "praeparat" in text and "nicht gefunden" in text:
        return "unknown_praeparat"
    if "ungueltiger typ" in text:
        return "invalid_type"
    if "charge" in text and "leer" in text:
        return "missing_charge"
    if "anzahl" in text:
        return "invalid_amount"
    if "datum" in text:
        return "invalid_date"
    if "verfall" in text:
        return "invalid_expiry"
    if "duplikat" in text:
        return "duplicate"
    return "other"


def _summarize_import_errors(errors: list[str]) -> dict[str, Any]:
    by_code: dict[str, int] = {}
    for item in errors:
        code = _classify_import_error_code(item)
        by_code[code] = by_code.get(code, 0) + 1
    return {"total": int(len(errors)), "by_code": by_code}


def _build_import_template(repository: SqliteRepository) -> bytes:
    _ensure_import_dependencies()
    wb = Workbook()
    ws = wb.active
    ws.title = "Bewegungen"
    depots = [row["name"] for row in repository.list_depots(limit=10000, offset=0, q="")]
    praeparate = [row["name"] for row in repository.list_praeparate(limit=10000, offset=0, q="")]
    typen = ["Zugang", "Abgang", "Vernichtung"]

    validation = wb.create_sheet("Validierung")
    for idx, name in enumerate(depots, start=1):
        validation.cell(row=idx, column=1, value=name)
    for idx, name in enumerate(praeparate, start=1):
        validation.cell(row=idx, column=2, value=name)
    for idx, name in enumerate(typen, start=1):
        validation.cell(row=idx, column=3, value=name)
    validation.sheet_state = "hidden"

    ws.merge_cells("A1:H1")
    ws["A1"] = "Depot in B2 waehlen, danach Daten ab Zeile 4 eintragen."
    ws["A2"] = "Depot:"
    ws["B2"] = depots[0] if depots else ""
    headers = ["Depot", "Praeparat", "Typ", "Charge", "Verfall", "Datum", "Anzahl", "Empfaenger"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=header)
    for row_idx in range(4, 104):
        ws.cell(row=row_idx, column=1, value=f'=IF(B{row_idx}<>"",$B$2,"")')

    if depots:
        dv_depot = DataValidation(type="list", formula1=f"Validierung!$A$1:$A${len(depots)}", allow_blank=False)
        ws.add_data_validation(dv_depot)
        dv_depot.add("B2")
    if praeparate:
        dv_praep = DataValidation(type="list", formula1=f"Validierung!$B$1:$B${len(praeparate)}", allow_blank=False)
        ws.add_data_validation(dv_praep)
        dv_praep.add("B4:B103")
    dv_typ = DataValidation(type="list", formula1=f"Validierung!$C$1:$C${len(typen)}", allow_blank=False)
    ws.add_data_validation(dv_typ)
    dv_typ.add("C4:C103")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _csv_response(filename: str, headers: list[str], rows: list[list[object]]) -> StreamingResponse:
    text_io = io.StringIO()
    writer = csv.writer(text_io)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    buffer = BytesIO(text_io.getvalue().encode("utf-8"))
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pdf_response(filename: str, title: str, headers: list[str], rows: list[list[object]], kpis: dict) -> StreamingResponse:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PDF-Export erfordert reportlab.") from exc

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Heading1"]), Spacer(1, 8)]
    if kpis:
        for key, value in kpis.items():
            elements.append(Paragraph(f"<b>{key}:</b> {value}", styles["Normal"]))
        elements.append(Spacer(1, 8))
    table_data = [headers] + rows
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#007aff")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pptx_response(filename: str, title: str, headers: list[str], rows: list[list[object]], kpis: dict) -> StreamingResponse:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PPT-Export erfordert python-pptx.") from exc

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_box = title_slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.size = Pt(24)
    title_tf.paragraphs[0].font.bold = True

    y = 1.0
    if kpis:
        kpi_box = title_slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(8.8), Inches(1.4))
        kpi_tf = kpi_box.text_frame
        first = True
        for key, value in kpis.items():
            p = kpi_tf.paragraphs[0] if first else kpi_tf.add_paragraph()
            first = False
            p.text = f"{key}: {value}"
            p.font.size = Pt(14)
        y += 1.5

    max_rows = min(len(rows), 18)
    cols = len(headers)
    table_shape = title_slide.shapes.add_table(cols=cols, rows=max_rows + 1, left=Inches(0.4), top=Inches(y), width=Inches(9.0), height=Inches(5.0))
    table = table_shape.table
    for col_idx, header in enumerate(headers):
        table.cell(0, col_idx).text = str(header)
    for r in range(max_rows):
        for c in range(cols):
            table.cell(r + 1, c).text = str(rows[r][c])

    with tempfile.NamedTemporaryFile(suffix=".pptx") as tmp:
        prs.save(tmp.name)
        tmp.seek(0)
        payload = tmp.read()
    out = BytesIO(payload)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


ALLOWED_USER_ROLES = {"Admin", "User"}
PERMISSION_DEFINITIONS = [
    {"key": "masterdata_read", "label": "Stammdaten lesen"},
    {"key": "masterdata_write", "label": "Stammdaten bearbeiten"},
    {"key": "movements_read", "label": "Bewegungen lesen"},
    {"key": "movements_write", "label": "Bewegungen erfassen"},
    {"key": "settings_read", "label": "Grundeinstellungen lesen"},
    {"key": "settings_write", "label": "Grundeinstellungen bearbeiten"},
    {"key": "import_use", "label": "Import nutzen"},
    {"key": "email_use", "label": "E-Mail nutzen"},
    {"key": "reports_view", "label": "Auswertungen sehen"},
    {"key": "audit_view", "label": "Audit sehen"},
    {"key": "users_manage", "label": "Benutzer verwalten"},
    {"key": "backup_manage", "label": "Backup verwalten"},
]
ALL_PERMISSION_KEYS = {item["key"] for item in PERMISSION_DEFINITIONS}
DEFAULT_USER_PERMISSIONS = {
    "masterdata_read",
    "movements_read",
    "movements_write",
    "settings_read",
    "import_use",
    "email_use",
    "reports_view",
}
PERMISSION_TEMPLATES = [
    {"key": "readonly", "label": "ReadOnly", "permissions": ["masterdata_read", "movements_read", "settings_read", "reports_view"]},
    {"key": "disponent", "label": "Disponent", "permissions": ["masterdata_read", "movements_read", "movements_write", "settings_read", "import_use", "email_use", "reports_view"]},
    {"key": "reporting", "label": "Reporting", "permissions": ["masterdata_read", "movements_read", "reports_view"]},
    {"key": "ops_admin", "label": "Ops Admin", "permissions": ["masterdata_read", "masterdata_write", "movements_read", "movements_write", "settings_read", "settings_write", "import_use", "email_use", "reports_view", "audit_view", "backup_manage"]},
]


def _parse_permission_list(raw_permissions: object) -> set[str]:
    if raw_permissions in (None, ""):
        return set()
    if isinstance(raw_permissions, str):
        try:
            parsed = json.loads(raw_permissions)
        except json.JSONDecodeError:
            return set()
    elif isinstance(raw_permissions, (list, tuple, set)):
        parsed = list(raw_permissions)
    else:
        return set()
    result: set[str] = set()
    for item in parsed:
        key = str(item or "").strip()
        if key in ALL_PERMISSION_KEYS:
            result.add(key)
    return result


def _permissions_for_role(role: str | None, raw_permissions: object) -> set[str]:
    if role == "Admin":
        return set(ALL_PERMISSION_KEYS)
    parsed = _parse_permission_list(raw_permissions)
    return parsed or set(DEFAULT_USER_PERMISSIONS)


def _permissions_json_for_storage(role: str, requested: list[str] | None) -> str:
    allowed = _permissions_for_role(role, requested)
    return json.dumps(sorted(allowed), ensure_ascii=True)


def _normalize_user_row(row: tuple) -> dict:
    effective_permissions = sorted(_permissions_for_role(str(row[2]), row[10] if len(row) > 10 else None))
    return {
        "id": int(row[0]),
        "username": str(row[1]),
        "role": str(row[2]),
        "email": row[3],
        "created_at": row[4],
        "last_login": row[5],
        "is_active": bool(row[6]),
        "failed_attempts": int(row[7] or 0),
        "locked_until": row[8],
        "is_default_password": bool(row[9]),
        "permissions": effective_permissions,
    }


def _get_user_flags(security: SecurityManager, username: str) -> dict[str, bool]:
    row = security.cur.execute(
        "SELECT is_default_password, role, permissions FROM users WHERE username = ?",
        ((username or "").strip(),),
    ).fetchone()
    if not row:
        return {"requires_password_change": False, "permissions": sorted(DEFAULT_USER_PERMISSIONS)}  # nosec B105: boolean flag, not a password
    return {
        "requires_password_change": bool(row[0]),
        "permissions": sorted(_permissions_for_role(str(row[1]), row[2])),
    }


def _get_avatar_path_for_user(security: SecurityManager, username: str) -> Path | None:
    row = security.cur.execute(
        "SELECT avatar_path FROM users WHERE username = ?",
        ((username or "").strip(),),
    ).fetchone()
    if not row or not row[0]:
        return None
    avatar_path = Path(str(row[0])).expanduser()
    if not avatar_path.exists() or not avatar_path.is_file():
        return None
    return avatar_path


def _safe_backup_label(value: str) -> str:
    return re.sub(r"[^a-z0-9_-]+", "_", (value or "").strip().lower()) or "manual"


def _normalize_json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _parse_iso_datetime(value: object) -> datetime | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    if raw.endswith("Z"):
        raw = raw[:-1] + "+00:00"
    try:
        parsed = datetime.fromisoformat(raw)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


def _create_backup_snapshot(source_db_path: Path, backups_dir: Path, label: str) -> Path:
    backups_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"ndhub_{_safe_backup_label(label)}_{timestamp}.db"
    target_path = backups_dir / filename
    with sqlite3.connect(str(source_db_path)) as source_conn:
        with sqlite3.connect(str(target_path)) as target_conn:
            source_conn.backup(target_conn)
    return target_path


def _create_mariadb_backup_snapshot(backups_dir: Path, label: str) -> Path:
    try:
        import pymysql
        from pymysql.cursors import DictCursor
    except ImportError as exc:
        raise RuntimeError("PyMySQL ist fuer MariaDB-Backups erforderlich.") from exc

    settings = resolve_mariadb_settings()
    backups_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"ndhub_{_safe_backup_label(label)}_{timestamp}.mariadb.json"
    target_path = backups_dir / filename

    payload: dict[str, object] = {
        "engine": "mariadb",
        "database": settings.database,
        "created_at": datetime.now(UTC).isoformat(),
        "tables": {},
    }
    conn = pymysql.connect(
        host=settings.host,
        port=settings.port,
        user=settings.user,
        password=settings.password,
        database=settings.database,
        charset="utf8mb4",
        autocommit=False,
        cursorclass=DictCursor,
    )
    try:
        with conn.cursor() as cur:
            cur.execute("SHOW TABLES")
            rows = cur.fetchall()
            table_names = [str(next(iter(row.values()))) for row in rows]
            table_payload: dict[str, list[dict]] = {}
            for table_name in sorted(table_names):
                quoted_table = _quote_sql_identifier(table_name)
                cur.execute("".join(["SELECT * FROM ", quoted_table]))
                data_rows = []
                for row in cur.fetchall():
                    normalized = {str(k): _normalize_json_value(v) for k, v in dict(row).items()}
                    data_rows.append(normalized)
                table_payload[table_name] = data_rows
            payload["tables"] = table_payload
    finally:
        conn.close()

    with target_path.open("w", encoding="utf-8") as out:
        json.dump(payload, out, ensure_ascii=False)
    return target_path


def _list_backup_files(backups_dir: Path, limit: int = 200) -> list[dict]:
    if not backups_dir.exists():
        return []
    rows: list[dict] = []
    backup_paths = sorted(backups_dir.glob("ndhub_*"), key=lambda item: item.stat().st_mtime, reverse=True)
    for path in backup_paths:
        if not path.is_file():
            continue
        if not (
            path.name.endswith(".db")
            or path.name.endswith(".sqlite")
            or path.name.endswith(".sqlite3")
            or path.name.endswith(".mariadb.json")
        ):
            continue
        stat = path.stat()
        rows.append(
            {
                "filename": path.name,
                "size_bytes": int(stat.st_size),
                "created_at": datetime.fromtimestamp(stat.st_mtime, tz=UTC).isoformat(),
                "kind": "mariadb_json" if path.name.endswith(".mariadb.json") else "sqlite_file",
            }
        )
        if len(rows) >= max(1, min(int(limit), 1000)):
            break
    return rows


def _run_auto_backup_if_due(source_db_path: Path, backups_dir: Path, interval_hours: int) -> Path | None:
    safe_interval_hours = max(1, min(int(interval_hours), 24 * 30))
    now_ts = datetime.now(UTC).timestamp()
    latest_auto_mtime = None
    for path in backups_dir.glob("ndhub_auto_*.db"):
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        if latest_auto_mtime is None or mtime > latest_auto_mtime:
            latest_auto_mtime = mtime
    if latest_auto_mtime is not None:
        elapsed_hours = (now_ts - latest_auto_mtime) / 3600
        if elapsed_hours < safe_interval_hours:
            return None
    return _create_backup_snapshot(source_db_path, backups_dir, "auto")


def _run_auto_backup_if_due_mariadb(backups_dir: Path, interval_hours: int) -> Path | None:
    safe_interval_hours = max(1, min(int(interval_hours), 24 * 30))
    now_ts = datetime.now(UTC).timestamp()
    latest_auto_mtime = None
    for path in backups_dir.glob("ndhub_auto_*.mariadb.json"):
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        if latest_auto_mtime is None or mtime > latest_auto_mtime:
            latest_auto_mtime = mtime
    if latest_auto_mtime is not None:
        elapsed_hours = (now_ts - latest_auto_mtime) / 3600
        if elapsed_hours < safe_interval_hours:
            return None
    return _create_mariadb_backup_snapshot(backups_dir, "auto")


def _restore_mariadb_backup_payload(payload: dict) -> None:
    engine = str(payload.get("engine") or "").strip().lower()
    if engine and engine != "mariadb":
        raise ValueError("Backup-Format ungueltig: engine ist nicht mariadb.")
    try:
        import pymysql
    except ImportError as exc:
        raise RuntimeError("PyMySQL ist fuer MariaDB-Restore erforderlich.") from exc

    tables = payload.get("tables")
    if not isinstance(tables, dict):
        raise ValueError("Backup-Format ungueltig: 'tables' fehlt.")
    required = {"users", "depots", "praeparate", "bewegungen"}
    present = {str(name) for name in tables.keys()}
    if not required.issubset(present):
        missing = ", ".join(sorted(required - present))
        raise ValueError(f"Backup ungueltig, fehlende Tabellen: {missing}")

    settings = resolve_mariadb_settings()
    conn = pymysql.connect(
        host=settings.host,
        port=settings.port,
        user=settings.user,
        password=settings.password,
        database=settings.database,
        charset="utf8mb4",
        autocommit=False,
    )
    preferred_order = [
        "users",
        "depots",
        "praeparate",
        "depot_praeparate",
        "kontakte",
        "bewegungen",
        "email_verlauf",
        "user_activity_log",
        "api_audit_log",
    ]
    ordered_tables = [name for name in preferred_order if name in tables]
    ordered_tables.extend(name for name in tables.keys() if name not in ordered_tables)

    allowed_restore_tables = set(preferred_order)
    invalid_tables = [name for name in ordered_tables if name not in allowed_restore_tables]
    if invalid_tables:
        raise ValueError(f"Backup enthaelt ungueltige Tabellen: {', '.join(sorted(invalid_tables))}")

    try:
        with conn.cursor() as cur:
            cur.execute("SET FOREIGN_KEY_CHECKS=0")
            for table_name in ordered_tables:
                quoted_table = _quote_sql_identifier(table_name)
                cur.execute("".join(["TRUNCATE TABLE ", quoted_table]))
            for table_name in ordered_tables:
                table_rows = tables.get(table_name) or []
                if not isinstance(table_rows, list) or not table_rows:
                    continue
                first = table_rows[0]
                if not isinstance(first, dict):
                    raise ValueError(f"Backup-Format ungueltig in Tabelle {table_name}.")
                columns = list(first.keys())
                for col in columns:
                    _quote_sql_identifier(col)
                col_sql = ", ".join(_quote_sql_identifier(col) for col in columns)
                placeholders = ", ".join(["%s"] * len(columns))
                query = "".join(
                    [
                        "INSERT INTO ",
                        _quote_sql_identifier(table_name),
                        " (",
                        col_sql,
                        ") VALUES (",
                        placeholders,
                        ")",
                    ]
                )
                values = []
                for row in table_rows:
                    if not isinstance(row, dict):
                        continue
                    values.append(tuple(row.get(col) for col in columns))
                if values:
                    cur.executemany(query, values)
            cur.execute("SET FOREIGN_KEY_CHECKS=1")
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _read_upload_size_bytes(upload_file: UploadFile) -> int:
    upload_file.file.seek(0, io.SEEK_END)
    size_bytes = int(upload_file.file.tell() or 0)
    upload_file.file.seek(0)
    return max(0, size_bytes)


def _validate_restore_upload_size(upload_file: UploadFile, max_bytes: int) -> int:
    size_bytes = _read_upload_size_bytes(upload_file)
    if size_bytes <= 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Backup-Datei ist leer.")
    if size_bytes > max_bytes:
        max_mb = round(max_bytes / (1024 * 1024), 2)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Backup-Datei ist zu gross (max. {max_mb} MB).",
        )
    return size_bytes


def _is_valid_sqlite_file(path: Path) -> bool:
    try:
        with path.open("rb") as handle:
            header = handle.read(16)
    except OSError:
        return False
    return header.startswith(b"SQLite format 3")


def create_app(db_path: str | None = None) -> FastAPI:
    """Application factory for runtime and tests."""
    db_engine = resolve_db_engine(default="sqlite")
    runtime_paths = resolve_runtime_paths(db_path=db_path)
    database_path = str(runtime_paths.database_path)
    source_db_path = runtime_paths.database_path
    attachments_dir = runtime_paths.attachments_dir
    backups_dir = runtime_paths.backups_dir
    auto_backup_hours = resolve_auto_backup_hours(default=24)
    max_backup_restore_mb = resolve_max_backup_restore_mb(default=200)
    max_backup_restore_bytes = max_backup_restore_mb * 1024 * 1024
    email_delivery_settings = resolve_email_delivery_settings(default_mode="draft")
    feature_multi_institution = (os.environ.get("ND_HUB_FEATURE_MULTI_INSTITUTION", "1") or "1").strip().lower() in {"1", "true", "yes", "on"}
    feature_institution_map = (os.environ.get("ND_HUB_FEATURE_INSTITUTION_MAP", "1") or "1").strip().lower() in {"1", "true", "yes", "on"}
    sqlite_fallback_repository = SqliteRepository(database_path)
    if db_engine == "mariadb":
        repository = MariaDbRepository(
            settings=resolve_mariadb_settings(),
            fallback=sqlite_fallback_repository,
            dual_write_sqlite=resolve_dual_write_sqlite(default=True),
        )
    else:
        repository = sqlite_fallback_repository
    security = SecurityManager(database_path)
    token_store = TokenStore(storage_path=database_path)

    def _mask_token(token: str) -> str:
        token = str(token or "")
        if len(token) <= 10:
            return "*" * len(token)
        return f"{token[:4]}{'*' * (len(token) - 8)}{token[-4:]}"

    def _resolve_archive_status(expires_at: datetime | None, revoked_at: datetime | None, token: str) -> str:
        if isinstance(expires_at, datetime) and expires_at <= datetime.now(UTC):
            return "abgelaufen"
        if revoked_at is not None:
            return "widerrufen"
        return "aktiv" if token_store.get(token) is not None else "widerrufen"

    def _user_permissions(username: str, role: str | None) -> set[str]:
        if role == "Admin":
            return set(ALL_PERMISSION_KEYS)
        row = security.cur.execute(
            "SELECT permissions FROM users WHERE username = ?",
            ((username or "").strip(),),
        ).fetchone()
        raw_permissions = row[0] if row else None
        return _permissions_for_role(role, raw_permissions)

    def _allowed_depot_permissions(session: SessionInfo) -> dict[int, dict[str, bool]]:
        if session.role == "Admin":
            depots = repository.list_depots(limit=2000, offset=0)
            return {
                int(item["id"]): {"can_read": True, "can_write": True}
                for item in depots
                if int(item.get("id", 0)) > 0
            }
        if not hasattr(repository, "list_user_depot_permissions"):
            return {}
        rows = repository.list_user_depot_permissions(session.username)
        result: dict[int, dict[str, bool]] = {}
        for row in rows:
            depot_id = int(row.get("depot_id") or 0)
            if depot_id <= 0:
                continue
            result[depot_id] = {
                "can_read": bool(int(row.get("can_read") or 0)),
                "can_write": bool(int(row.get("can_write") or 0)),
            }
        return result

    def _ensure_depot_access(session: SessionInfo, depot_id: int, require_write: bool = False) -> None:
        if session.role == "Admin":
            return
        allowed = _allowed_depot_permissions(session)
        depot_rights = allowed.get(int(depot_id))
        if not depot_rights:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Depot nicht freigegeben.")
        key = "can_write" if require_write else "can_read"
        if not bool(depot_rights.get(key)):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Depot-Berechtigung fehlt.")

    def _allowed_depot_ids(session: SessionInfo, require_write: bool = False) -> list[int]:
        if session.role == "Admin":
            rows = repository.list_depots(limit=5000, offset=0)
            return [int(row["id"]) for row in rows if int(row.get("id", 0)) > 0]
        key = "can_write" if require_write else "can_read"
        allowed = _allowed_depot_permissions(session)
        return sorted([depot_id for depot_id, rights in allowed.items() if rights.get(key)])

    def _scoped_requested_ids(session: SessionInfo, requested_ids: list[int], require_write: bool = False) -> list[int]:
        allowed_ids = set(_allowed_depot_ids(session, require_write=require_write))
        if session.role != "Admin" and not allowed_ids:
            return []
        if not requested_ids:
            return sorted(allowed_ids) if session.role != "Admin" else []
        safe_requested = sorted({int(item) for item in requested_ids if int(item) > 0})
        if session.role == "Admin":
            return safe_requested
        forbidden = [item for item in safe_requested if item not in allowed_ids]
        if forbidden:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Mindestens ein Depot ist nicht freigegeben.")
        return safe_requested

    def require_permission(permission_key: str):
        def _dependency(session: SessionInfo = Depends(get_current_session)) -> SessionInfo:
            if permission_key not in ALL_PERMISSION_KEYS:
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail=f"Unbekannte Berechtigung: {permission_key}",
                )
            if session.role == "Admin":
                return session
            permissions = _user_permissions(session.username, session.role)
            if permission_key not in permissions:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Berechtigung fehlt.",
                )
            return session

        return _dependency

    @asynccontextmanager
    async def _lifespan(_app: FastAPI):
        try:
            try:
                if db_engine == "mariadb":
                    created = _run_auto_backup_if_due_mariadb(backups_dir, auto_backup_hours)
                else:
                    created = _run_auto_backup_if_due(source_db_path, backups_dir, auto_backup_hours)
                if created is not None:
                    repository.log_audit(
                        username="system",
                        action="create",
                        resource_type="backup",
                        details={"operation": "auto", "filename": created.name},
                    )
            except Exception:
                # Auto backup should never block startup.
                logger.warning("Auto-backup during startup failed", exc_info=True)
            yield
        finally:
            security.close()
            token_store.close()

    app = FastAPI(title="ND-Hub Backend", version="0.1.1", lifespan=_lifespan)
    app.state.repository = repository
    app.state.security = security
    app.state.token_store = token_store
    app.state.database_path = database_path
    app.state.web_dir = Path(__file__).parent / "web"
    app.state.attachments_dir = attachments_dir
    app.state.backups_dir = backups_dir
    app.state.auto_backup_hours = auto_backup_hours
    app.state.max_backup_restore_mb = max_backup_restore_mb
    app.state.db_engine = db_engine
    app.state.email_delivery_settings = email_delivery_settings
    app.state.feature_multi_institution = feature_multi_institution
    app.state.feature_institution_map = feature_institution_map

    # ---- Security-Middleware (Fix für #31) ----
    # 1) TrustedHost: blockiert Requests mit unbekanntem Host-Header
    #    (verhindert Host-Header-Injection / Cache-Poisoning).
    #    Default ["*"] = alle Hosts; produktiv bitte ND_HUB_ALLOWED_HOSTS setzen.
    _allowed_hosts_raw = os.environ.get("ND_HUB_ALLOWED_HOSTS", "*").strip()
    if _allowed_hosts_raw == "*":
        _allowed_hosts = ["*"]
    else:
        _allowed_hosts = [h.strip() for h in _allowed_hosts_raw.split(",") if h.strip()]
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=_allowed_hosts)

    # 2) CORS: restriktiver Default (allow_origins=[]). Cross-Origin nur
    #    explizit via ND_HUB_CORS_ORIGINS (Komma-getrennte Allowlist).
    _cors_raw = os.environ.get("ND_HUB_CORS_ORIGINS", "").strip()
    _cors_origins = [o.strip() for o in _cors_raw.split(",") if o.strip()] if _cors_raw else []
    app.add_middleware(
        CORSMiddleware,
        allow_origins=_cors_origins,
        allow_credentials=True,
        allow_methods=["GET", "POST", "PUT", "DELETE", "PATCH", "OPTIONS"],
        allow_headers=["Authorization", "Content-Type", "Accept"],
    )

    # 3) Security-Header-Middleware: schraubt Standard-Schutz-Header auf
    #    jede Response. HSTS nur bei HTTPS-Request, damit Dev über HTTP ok bleibt.
    @app.middleware("http")
    async def _security_headers_middleware(request: Request, call_next):
        response = await call_next(request)
        response.headers.setdefault("X-Content-Type-Options", "nosniff")
        response.headers.setdefault("X-Frame-Options", "DENY")
        response.headers.setdefault("Referrer-Policy", "no-referrer")
        if request.url.scheme == "https":
            response.headers.setdefault(
                "Strict-Transport-Security",
                "max-age=63072000; includeSubDomains",
            )
        return response
    # ---- /Security-Middleware ----

    # ---- Login-Rate-Limit (Fix für #32) ----
    # Per-IP via slowapi (ENV-konfigurierbar; Default 5/min).
    # Per-Username-Lockout: 10 Fehlversuche -> 15 min Sperre (ENV-konfigurierbar).
    # Audit-Log-Eintrag bei Lockout.
    _login_ip_limit = os.environ.get("ND_HUB_LOGIN_IP_LIMIT", "5/minute")
    _login_max_fails = int(os.environ.get("ND_HUB_LOGIN_MAX_FAILS", "10"))
    _login_lockout_seconds = int(os.environ.get("ND_HUB_LOGIN_LOCKOUT_SECONDS", "900"))

    login_limiter = Limiter(key_func=get_remote_address, default_limits=[_login_ip_limit])
    app.state.login_limiter = login_limiter

    # Per-Username-Lockout-Tracker auf app.state (statt closure-lokal),
    # damit er pro App-Instanz eindeutig ist und von Tests verifiziert werden kann.
    # Struktur: {username: (fails, lockout_until_epoch)}
    # In-Memory (Single-Instance). Thread-safe via Lock.
    import threading as _threading
    app.state.login_lockout_state = {}
    app.state.login_lockout_lock = _threading.Lock()
    app.state.login_max_fails = _login_max_fails
    app.state.login_lockout_seconds = _login_lockout_seconds

    def _check_login_lockout(username: str) -> int | None:
        """Returnt verbleibende Sekunden, wenn locked; sonst None."""
        with app.state.login_lockout_lock:
            entry = app.state.login_lockout_state.get(username)
            if not entry:
                return None
            fails, lockout_until = entry
            # Eintraege mit unendlicher Ablaufzeit sind reine Counter
            # (noch nicht max_fails erreicht) und sperren nicht.
            if lockout_until == float("inf"):
                return None
            now = datetime.now().timestamp()
            if lockout_until > now:
                return int(lockout_until - now)
            # Lockout abgelaufen — Counter zuruecksetzen
            app.state.login_lockout_state.pop(username, None)
            return None

    def _record_login_failure(username: str) -> int | None:
        """Verzeichnet einen Fehlversuch. Returnt lockout-Restsekunden, wenn Lockout ausgeloest."""
        with app.state.login_lockout_lock:
            state = app.state.login_lockout_state
            entry = state.get(username)
            if entry is None:
                fails, lockout_until = 0, 0.0
            else:
                fails, lockout_until = entry
            fails += 1
            if fails >= app.state.login_max_fails:
                lockout_until = datetime.now().timestamp() + app.state.login_lockout_seconds
                state[username] = (fails, lockout_until)
                return app.state.login_lockout_seconds
            # Counter-Eintrag ohne aktiven Lockout: Ablaufzeit auf 'unendlich'
            # setzen, damit _check_login_lockout ihn nicht als 'abgelaufen'
            # verwirft, bevor max_fails erreicht ist.
            state[username] = (fails, float("inf"))
            return None

    def _clear_login_lockout(username: str) -> None:
        with app.state.login_lockout_lock:
            app.state.login_lockout_state.pop(username, None)

    def _audit_login_lockout(username: str, ip: str) -> None:
        """Audit-Log-Eintrag bei Lockout. Best-effort, kein Raise bei Fehler."""
        try:
            security.log_activity(
                user_id=0,
                username=username,
                action="login_lockout",
                details=f"ip={ip} max_fails={app.state.login_max_fails} lockout_seconds={app.state.login_lockout_seconds}",
                ip=ip,
            )
        except Exception:
            logger.warning("Audit-Log fuer login_lockout fehlgeschlagen", exc_info=True)

    # slowapi-Exception-Handler: 429 + Retry-After-Header bei IP-Limit
    async def _rate_limit_handler(_request: _StarletteRequest, exc: RateLimitExceeded):
        # Retry-After aus slowapi-Limit-Strategie ableiten
        retry_after = 60  # Default-Fallback
        try:
            # slowapi speichert Limit in exc.limit, Detail in exc.detail
            if hasattr(exc, "limit") and exc.limit and hasattr(exc.limit, "seconds"):
                retry_after = exc.limit.seconds
        except Exception:
            logger.exception("Unexpected error in endpoint")
        return _JSONResponse(
            status_code=429,
            content={"detail": f"Zu viele Login-Versuche. Bitte {retry_after}s warten."},
            headers={"Retry-After": str(retry_after)},
        )

    app.add_exception_handler(RateLimitExceeded, _rate_limit_handler)
    # ---- /Login-Rate-Limit ----

    app.mount(
        "/web",
        StaticFiles(directory=app.state.web_dir, html=False),
        name="web",
    )

    @app.get("/")
    def web_index() -> FileResponse:
        return FileResponse(app.state.web_dir / "index.html")

    @app.get("/health")
    def health() -> dict[str, str]:
        return {"status": "ok", "db_engine": str(app.state.db_engine)}

    @app.post("/auth/login", response_model=LoginResponse)
    @login_limiter.limit(_login_ip_limit)
    def login(payload: LoginRequest, request: Request) -> LoginResponse:
        # Fix #32: per-Username-Lockout pruefen BEVOR authenticate() laeuft
        # (verhindert, dass PBKDF2-CPU-Zeit verschwendet wird, waehrend locked)
        username_attempt = payload.username.strip()
        remaining_lockout = _check_login_lockout(username_attempt)
        if remaining_lockout is not None:
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail=f"Account temporaer gesperrt. Bitte {remaining_lockout}s warten.",
                headers={"Retry-After": str(remaining_lockout)},
            )
        ok, _message = security.authenticate(username_attempt, payload.password)
        if not ok:
            lockout_seconds = _record_login_failure(username_attempt)
            if lockout_seconds is not None:
                # Lockout frisch ausgeloest
                _audit_login_lockout(username_attempt, request.client.host if request.client else "unknown")
                raise HTTPException(
                    status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                    detail=f"Account gesperrt nach {app.state.login_max_fails} Fehlversuchen. "
                           f"Bitte {lockout_seconds}s warten.",
                    headers={"Retry-After": str(lockout_seconds)},
                )
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Benutzername oder Passwort ungueltig.",
            )
        # Erfolgreich: Lockout-Counter zuruecksetzen
        _clear_login_lockout(username_attempt)
        username = security.get_current_user() or payload.username
        user_flags = _get_user_flags(security, username)
        token = token_store.issue(
            username=username,
            role=security.get_current_role(),
        )
        return LoginResponse(
            token=token,
            username=username,
            role=security.get_current_role(),
            requires_password_change=user_flags["requires_password_change"],
            permissions=user_flags["permissions"],
        )


    # ---- Web-only domain routers (Issue #60) ----
    from backend.routers.desktop_sync_auth import create_desktop_sync_auth_router
    from backend.routers.institutions import create_institutions_router
    from backend.routers.sync import create_sync_router

    app.include_router(
        create_desktop_sync_auth_router(
            token_store=token_store,
            get_current_session=get_current_session,
            DesktopSyncTokenResponse=DesktopSyncTokenResponse,
            DesktopSyncTokenCreateRequest=DesktopSyncTokenCreateRequest,
            DesktopSyncTokenArchiveItem=DesktopSyncTokenArchiveItem,
            DesktopSyncTokenRevokeRequest=DesktopSyncTokenRevokeRequest,
            _parse_iso_datetime=_parse_iso_datetime,
            _mask_token=_mask_token,
            _resolve_archive_status=_resolve_archive_status,
        )
    )
    app.include_router(
        create_institutions_router(
            app=app,
            repository=repository,
            require_permission=require_permission,
            get_current_session=get_current_session,
            OnboardingInstitutionSetupRequest=OnboardingInstitutionSetupRequest,
            InstitutionUpsertRequest=InstitutionUpsertRequest,
            _allowed_depot_ids=_allowed_depot_ids,
            _require_http_scheme=_require_http_scheme,
        )
    )
    app.include_router(
        create_sync_router(
            repository=repository,
            require_permission=require_permission,
            get_current_session=get_current_session,
            SyncPushRequest=SyncPushRequest,
            SyncPushChangeItem=SyncPushChangeItem,
            SyncPullRequest=SyncPullRequest,
            _allowed_depot_ids=_allowed_depot_ids,
            _user_permissions=_user_permissions,
        )
    )
    # ---- /Web-only domain routers ----

    # ---- Shared routers (Issues #60/#61/#65) ----
    from backend.routers.auth import create_auth_router
    from backend.routers.users import create_users_router
    from backend.routers.depots import create_depots_router

    def _enrich_auth_me(session: SessionInfo) -> dict:
        depot_permissions = _allowed_depot_permissions(session)
        return {
            "depot_permissions": [
                {"depot_id": depot_id, **rights}
                for depot_id, rights in sorted(depot_permissions.items())
            ],
        }

    _auth_router = create_auth_router(
        security=security,
        token_store=token_store,
        repository=repository,
        get_current_session=get_current_session,
        bearer_scheme=bearer_scheme,
        password_change_model=PasswordChangeRequest,
        get_user_flags=_get_user_flags,
        get_avatar_path_for_user=_get_avatar_path_for_user,
        enrich_me=_enrich_auth_me,
    )
    app.include_router(_auth_router)

    _users_router = create_users_router(
        security=security,
        repository=repository,
        require_permission=require_permission,
        normalize_user_row=_normalize_user_row,
        permissions_json_for_storage=_permissions_json_for_storage,
        allowed_user_roles=ALLOWED_USER_ROLES,
        hash_password=SecurityManager.hash_password,
        user_create_model=UserCreateRequest,
        user_update_model=UserUpdateRequest,
        user_password_reset_model=UserPasswordResetRequest,
        include_depot_permissions=True,
        user_depot_permission_update_model=UserDepotPermissionUpdateRequest,
        is_multi_institution=lambda: bool(app.state.feature_multi_institution),
    )
    app.include_router(_users_router)

    _depots_router = create_depots_router(
        repository=repository,
        require_permission=require_permission,
        depot_upsert_model=DepotUpsertRequest,
        depot_assignments_update_model=DepotAssignmentsUpdateRequest,
        kontakt_upsert_model=KontaktUpsertRequest,
        ensure_depot_access=_ensure_depot_access,
        resolve_list_allowed_ids=lambda session: (
            list(_allowed_depot_permissions(session).keys()) if session.role != "Admin" else None
        ),
        include_geo_fields=True,
    )
    app.include_router(_depots_router)

    from backend.routers.praeparate import create_praeparate_router
    from backend.routers.kontakte import create_kontakte_router
    from backend.routers.audit import create_audit_router
    from backend.routers.permissions import create_permissions_router
    from backend.routers.bewegungen import create_bewegungen_router

    app.include_router(
        create_permissions_router(
            get_current_session,
            permission_definitions=PERMISSION_DEFINITIONS,
            default_user_permissions=DEFAULT_USER_PERMISSIONS,
            permission_templates=PERMISSION_TEMPLATES,
        )
    )
    app.include_router(
        create_praeparate_router(
            repository=repository,
            require_permission=require_permission,
            praeparat_upsert_model=PraeparatUpsertRequest,
            include_extended_fields=True,
        )
    )
    app.include_router(
        create_kontakte_router(
            repository=repository,
            require_permission=require_permission,
            kontakt_upsert_model=KontaktUpsertRequest,
        )
    )
    app.include_router(
        create_audit_router(
            repository=repository,
            require_permission=require_permission,
        )
    )
    app.include_router(
        create_bewegungen_router(
            repository=repository,
            require_permission=require_permission,
            bewegung_create_model=BewegungCreateRequest,
            ensure_depot_access=_ensure_depot_access,
            scoped_requested_ids=_scoped_requested_ids,
            include_advanced_filters=True,
            include_export=True,
            csv_response=_csv_response,
            attachments_dir=app.state.attachments_dir,
            build_attachment_target_path=_build_attachment_target_path,
            persist_pdf_upload=_persist_pdf_upload,
            pdf_mime_types=PDF_MIME_TYPES,
        )
    )


    from backend.routers.imports import create_imports_router
    from backend.routers.verfall import create_verfall_router
    from backend.routers.dashboard import create_dashboard_router
    from backend.routers.notifications import create_notifications_router

    app.include_router(
        create_imports_router(
            repository=repository,
            require_permission=require_permission,
            build_import_template=_build_import_template,
            load_import_dataframe=_load_import_dataframe,
            map_import_columns=_map_import_columns,
            analyze_import_dataframe=_analyze_import_dataframe,
            enable_web_features=True,
            summarize_import_errors=_summarize_import_errors,
            build_import_execution_fingerprint=_build_import_execution_fingerprint,
            allowed_depot_ids=_allowed_depot_ids,
        )
    )
    app.include_router(
        create_verfall_router(
            repository=repository,
            require_permission=require_permission,
            parse_id_list_csv=_parse_id_list_csv,
            normalize_verfall_thresholds=_normalize_verfall_thresholds,
            enrich_verfall_rows=_enrich_verfall_rows,
            csv_response=_csv_response,
            enable_depot_scope=True,
            allowed_depot_ids=_allowed_depot_ids,
            scoped_requested_ids=_scoped_requested_ids,
        )
    )
    app.include_router(
        create_dashboard_router(
            repository=repository,
            require_permission=require_permission,
            normalize_verfall_thresholds=_normalize_verfall_thresholds,
            enrich_verfall_rows=_enrich_verfall_rows,
            enable_depot_scope=True,
            allowed_depot_ids=_allowed_depot_ids,
        )
    )
    app.include_router(
        create_notifications_router(
            repository=repository,
            require_permission=require_permission,
            normalize_verfall_thresholds=_normalize_verfall_thresholds,
            enrich_verfall_rows=_enrich_verfall_rows,
            enable_depot_scope=True,
            allowed_depot_ids=_allowed_depot_ids,
        )
    )


    from backend.routers.emails import create_emails_router
    from backend.routers.admin_backup import create_admin_backup_router

    def _close_security():
        nonlocal security
        security.close()

    def _reopen_security():
        nonlocal security, token_store
        security = SecurityManager(database_path)
        app.state.security = security
        token_store = TokenStore()
        app.state.token_store = token_store

    app.include_router(
        create_emails_router(
            repository=repository,
            require_permission=require_permission,
            email_recipient_preview_model=EmailRecipientPreviewRequest,
            email_draft_create_model=EmailDraftCreateRequest,
            enable_web_features=True,
            email_delivery_status_update_model=EmailDeliveryStatusUpdateRequest,
            get_email_delivery_settings=lambda: app.state.email_delivery_settings,
            email_delivery_missing_config=_email_delivery_missing_config,
            send_email_via_smtp=lambda *args, **kwargs: _send_email_via_smtp(*args, **kwargs),
        )
    )
    app.include_router(
        create_admin_backup_router(
            repository=repository,
            require_permission=require_permission,
            backups_dir=backups_dir,
            source_db_path=source_db_path,
            database_path=database_path,
            list_backup_files=_list_backup_files,
            create_backup_snapshot=_create_backup_snapshot,
            auto_backup_hours=auto_backup_hours,
            close_security=_close_security,
            reopen_security=_reopen_security,
            enable_web_features=True,
            db_engine=db_engine,
            max_restore_size_mb=max_backup_restore_mb,
            max_backup_restore_bytes=max_backup_restore_bytes,
            create_mariadb_backup_snapshot=_create_mariadb_backup_snapshot,
            validate_restore_upload_size=_validate_restore_upload_size,
            restore_mariadb_backup_payload=_restore_mariadb_backup_payload,
            is_valid_sqlite_file=_is_valid_sqlite_file,
        )
    )


    from backend.routers.reports import create_reports_router

    app.include_router(
        create_reports_router(
            repository=repository,
            require_permission=require_permission,
            parse_id_list_csv=_parse_id_list_csv,
            csv_response=_csv_response,
            pdf_response=_pdf_response,
            pptx_response=_pptx_response,
            enable_depot_scope=True,
            enable_rich_filenames=True,
            scoped_requested_ids=_scoped_requested_ids,
            allowed_depot_ids=_allowed_depot_ids,
            validated_report_date_range=_validated_report_date_range,
            build_report_export_filename=_build_report_export_filename,
            add_months=_add_months,
        )
    )

    # ---- /Shared routers ----

    return app


app = create_app()

