from collections import defaultdict
from datetime import date, datetime

from dateutil.relativedelta import relativedelta

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None

class PPTGenerator:
    """Generates Apple-styled PowerPoint Management Presentations"""
    
    def __init__(self, db_manager):
        if Presentation is None:
            raise RuntimeError(
                "PowerPoint-Export nicht verfügbar: Paket 'python-pptx' fehlt. "
                "Installieren Sie es mit 'pip install python-pptx'."
            )
        self.db = db_manager
        self.prs = Presentation()
        
        # Apple Light Theme Colors (Clean for PPT)
        self.c_blue = RGBColor(0, 122, 255)
        self.c_green = RGBColor(52, 199, 89)
        self.c_red = RGBColor(255, 59, 48)
        self.c_orange = RGBColor(255, 149, 0)
        self.c_purple = RGBColor(175, 82, 222)
        self.c_text = RGBColor(28, 28, 30)
        self.c_text_sec = RGBColor(142, 142, 147)
        self.c_bg = RGBColor(242, 242, 247)
        
    def _add_title_slide(self, title, subtitle):
        slide_layout = self.prs.slide_layouts[0] # Title slide
        slide = self.prs.slides.add_slide(slide_layout)
        
        # Format Title
        title_shape = slide.shapes.title
        title_shape.text = title
        title_shape.text_frame.paragraphs[0].font.name = 'Helvetica Neue'
        title_shape.text_frame.paragraphs[0].font.size = Pt(44)
        title_shape.text_frame.paragraphs[0].font.bold = True
        title_shape.text_frame.paragraphs[0].font.color.rgb = self.c_blue
        
        # Format Subtitle
        if slide.placeholders[1]:
            sub_shape = slide.placeholders[1]
            sub_shape.text = f"{subtitle}\nErstellt am {datetime.now().strftime('%d.%m.%Y')}"
            for p in sub_shape.text_frame.paragraphs:
                p.font.name = 'Helvetica Neue'
                p.font.size = Pt(24)
                p.font.color.rgb = self.c_text_sec
                
    def _add_kpi_slide(self, depots):
        # Fetch KPI Data
        bestand_data = []
        if not depots:
            bestand_data = self.db.get_bestandsentwicklung()
        else:
            bestand_data = self.db.get_bestandsentwicklung(depot_ids=depots)
            
        sum_soll = sum(d[2] for d in bestand_data)
        sum_ist = sum(d[3] for d in bestand_data)
        luecken = sum(1 for d in bestand_data if d[3] < d[2])
        quote = (sum_ist / sum_soll * 100) if sum_soll > 0 else 0
        
        # Verfalls-KPI (Nächste 3 Monate)
        if depots:
            placeholders = ','.join('?' * len(depots))
            where_clause = "depot_id IN (" + placeholders + ")"
            params = list(depots)
        else:
            where_clause = "1=1"
            params = []

        sql = (  # nosec B608: where_clause is built from ? placeholders or literal 1=1
            "SELECT SUM(anzahl)\n"
            "FROM bewegungen\n"
            "WHERE " + where_clause + "\n"  # nosec B608: where_clause built from ? placeholders or literal 1=1
            "  AND typ = 'Zugang'\n"
            "  AND (ausgang_datum IS NULL OR ausgang_datum = '')\n"
            "  AND verfall IS NOT NULL\n"
            "  AND verfall != ''\n"
            "  AND verfall <= ?"
        )
        in_3_months = (date.today() + relativedelta(months=3)).strftime("%Y-%m-%d")
        res = self.db.cur.execute(sql, params + [in_3_months]).fetchone()
        gef_verfall = res[0] if res and res[0] else 0
        
        # Create Slide
        slide_layout = self.prs.slide_layouts[5] # Blank with title
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_shape = slide.shapes.title
        title_shape.text = "Management Summary (KPIs)"
        title_shape.text_frame.paragraphs[0].font.color.rgb = self.c_text
        title_shape.text_frame.paragraphs[0].font.name = 'Helvetica Neue'

        # Draw 4 KPI Boxes
        kpis = [
            ("Bestandsquote", f"{quote:.1f}%", self.c_blue),
            ("Gesamtbestand", f"{int(sum_ist)} EH", self.c_green),
            ("Kritische Lücken", f"{luecken}", self.c_red if luecken>0 else self.c_green),
            ("Gefährdeter Verfall", f"{int(gef_verfall)} EH", self.c_orange)
        ]
        
        left = Inches(1)
        top = Inches(2.5)
        width = Inches(3.5)
        height = Inches(1.8)
        
        for i, (label, value, color) in enumerate(kpis):
            col = i % 2
            row = i // 2
            x = left + Inches(col * 4)
            y = top + Inches(row * 2.2)
            
            # Draw Box
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(248, 249, 250)
            shape.line.color.rgb = color
            shape.line.width = Pt(2)
            
            # Value Text
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = value
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = color
            p.font.name = 'Helvetica Neue'
            
            # Label Text
            p2 = tf.add_paragraph()
            p2.text = label
            p2.alignment = PP_ALIGN.CENTER
            p2.font.size = Pt(16)
            p2.font.color.rgb = self.c_text_sec
            p2.font.name = 'Helvetica Neue'

    def _add_bestand_chart(self, depots):
        # Data
        data = []
        if not depots:
            data = self.db.get_bestandsentwicklung()
        else:
            data = self.db.get_bestandsentwicklung(depot_ids=depots)
            
        if not data: return
        
        # Aggregate by Depot for cleaner chart
        depot_sums = defaultdict(lambda: {"soll": 0, "ist": 0})
        for depot_name, praep_name, soll, ist, diff in data:
            depot_sums[depot_name]["soll"] += soll
            depot_sums[depot_name]["ist"] += ist
            
        chart_data = ChartData()
        chart_data.categories = list(depot_sums.keys())
        chart_data.add_series('Soll-Bestand', [d["soll"] for d in depot_sums.values()])
        chart_data.add_series('Ist-Bestand', [d["ist"] for d in depot_sums.values()])

        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        title = slide.shapes.title
        title.text = "Soll-Ist-Abgleich der Depots"
        title.text_frame.paragraphs[0].font.color.rgb = self.c_text
        title.text_frame.paragraphs[0].font.name = 'Helvetica Neue'

        x, y, cx, cy = Inches(0.5), Inches(2), Inches(9), Inches(5)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
        ).chart
        
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        
        # Colors
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = self.c_text_sec
        chart.series[1].format.fill.solid()
        chart.series[1].format.fill.fore_color.rgb = self.c_blue
        
    def _add_action_items(self, depots):
        # Analyze data
        data = []
        if not depots:
            data = self.db.get_bestandsentwicklung()
        else:
            data = self.db.get_bestandsentwicklung(depot_ids=depots)
            
        kritische_luecken = [d for d in data if (d[2] - d[3]) > (d[2]*0.3) and d[2] > 0] # > 30% gap
        ueberschuss = [d for d in data if d[3] > d[2]*1.2 and d[2] > 0] # > 20% over
        
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1]) # Title and Content
        title = slide.shapes.title
        title.text = "KI-Analysen & Handlungsempfehlungen"
        title.text_frame.paragraphs[0].font.color.rgb = self.c_text
        title.text_frame.paragraphs[0].font.name = 'Helvetica Neue'
        
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        if not kritische_luecken and not ueberschuss:
            p = tf.paragraphs[0]
            p.text = "✅ Alle Bestände sind innerhalb optimaler Parameter."
            p.font.color.rgb = self.c_green
            return
            
        def add_bullet(text, color):
            p = tf.add_paragraph()
            p.text = text
            p.font.size = Pt(18)
            p.font.color.rgb = color
            p.font.name = 'Helvetica Neue'
            
        if kritische_luecken:
            add_bullet("⚠️ Kritische Unterbestände identifiziert:", self.c_red)
            for d in sorted(kritische_luecken, key=lambda x: x[2]-x[3], reverse=True)[:5]:
                add_bullet(f"   • {d[0]}: {d[1]} (Ist: {d[3]}, Soll: {d[2]}) -> {d[2]-d[3]} Einheiten fehlen.", self.c_text)
                
        if ueberschuss:
            add_bullet("\n💡 Signifikanter Überbestand (Ressourcenbindung):", self.c_blue)
            for d in sorted(ueberschuss, key=lambda x: x[3]-x[2], reverse=True)[:5]:
                add_bullet(f"   • {d[0]}: {d[1]} (Ist: {d[3]}, Soll: {d[2]}) -> +{d[3]-d[2]} Einheiten.", self.c_text)
                
    def _add_ranking_slide(self, depots):
        # Top Praeparate overall
        data = self.db.get_praeparat_ranking(depot_ids=depots if depots else None, limit=5)
        if not data: return
        
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        title = slide.shapes.title
        title.text = "Top 5 Präparate (nach Verbrauch)"
        title.text_frame.paragraphs[0].font.color.rgb = self.c_text
        title.text_frame.paragraphs[0].font.name = 'Helvetica Neue'
        
        chart_data = CategoryChartData()
        chart_data.categories = [d[0] for d in data]
        chart_data.add_series('Verbrauch', [d[2] for d in data])

        x, y, cx, cy = Inches(0.5), Inches(2), Inches(9), Inches(4.5)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED, x, y, cx, cy, chart_data
        ).chart
        
        chart.has_legend = False
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = self.c_purple

    def generate(self, filepath, config):
        """
        config = {
            'title': str,
            'subtitle': str,
            'depots': [id1, id2],
            'kpi': bool,
            'insights': bool,
            'ranking': bool,
            'bestand': bool
        }
        """
        self._add_title_slide(config.get('title', 'Jahresbilanz'), config.get('subtitle', ''))
        
        depots = config.get('depots', [])
        
        if config.get('kpi', True):
            self._add_kpi_slide(depots)
            
        if config.get('bestand', True):
            self._add_bestand_chart(depots)
            
        if config.get('ranking', True):
            self._add_ranking_slide(depots)
            
        if config.get('insights', True):
            self._add_action_items(depots)
            
        self.prs.save(filepath)
        return True
