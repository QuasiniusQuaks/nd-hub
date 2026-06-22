"""FastAPI app for ND-Hub MVP backend."""

from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import shutil
import sqlite3
import tempfile
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

try:
    import pandas as pd
except ImportError:  # pragma: no cover - optional runtime dependency
    pd = None
try:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - optional runtime dependency
    Workbook = None
    DataValidation = None

from backend.auth import (
    SessionInfo,
    TokenStore,
    bearer_scheme,
    get_current_session,
)
from backend.config import resolve_db_path
from backend.database import SqliteRepository
from security_manager import SecurityManager

logger = logging.getLogger(__name__)


class LoginRequest(BaseModel):
    username: str = Field(min_length=1)
    password: str = Field(min_length=1)


class LoginResponse(BaseModel):
    token: str
    username: str
    role: str | None
    requires_password_change: bool = False
    permissions: list[str] = Field(default_factory=list)


class PasswordChangeRequest(BaseModel):
    old_password: str = Field(min_length=1)
    new_password: str = Field(min_length=8)


class UserCreateRequest(BaseModel):
    username: str = Field(min_length=3)
    password: str = Field(min_length=8)
    role: str = Field(min_length=1)
    email: str | None = None
    is_active: bool = True
    permissions: list[str] | None = None


class UserUpdateRequest(BaseModel):
    username: str | None = Field(default=None, min_length=3)
    role: str | None = None
    email: str | None = None
    is_active: bool | None = None
    permissions: list[str] | None = None


class UserPasswordResetRequest(BaseModel):
    new_password: str = Field(min_length=8)


class BewegungCreateRequest(BaseModel):
    depot_id: int
    praeparat_id: int
    typ: str
    charge: str
    verfall: date
    datum: date
    anzahl: int = Field(gt=0)
    empfaenger: str | None = None


class DepotUpsertRequest(BaseModel):
    name: str = Field(min_length=1)
    adresse: str | None = None
    telefon: str | None = None
    email: str | None = None


class PraeparatUpsertRequest(BaseModel):
    name: str = Field(min_length=1)


class DepotAssignmentItem(BaseModel):
    praeparat_id: int
    sollbestand: int = Field(ge=0, default=0)


class DepotAssignmentsUpdateRequest(BaseModel):
    assignments: list[DepotAssignmentItem]


class KontaktUpsertRequest(BaseModel):
    name: str = Field(min_length=1)
    rolle: str | None = None
    telefon: str | None = None
    email: str | None = None


class EmailRecipientPreviewRequest(BaseModel):
    depot_ids: list[int] = Field(default_factory=list)


class EmailDraftCreateRequest(BaseModel):
    depot_ids: list[int] = Field(default_factory=list)
    betreff: str = Field(min_length=1)
    nachricht: str = ""


MAX_ATTACHMENT_SIZE_BYTES = 10 * 1024 * 1024
PDF_MIME_TYPES = {"application/pdf", "application/x-pdf"}
ALLOWED_IMPORT_TYPES = {"Zugang", "Abgang", "Vernichtung"}
IMPORT_REQUIRED_COLUMNS = ["depot", "praeparat", "typ", "charge", "verfall", "datum", "anzahl"]
IMPORT_COLUMN_ALIASES = {
    "depot": {"depot"},
    "praeparat": {"praeparat", "präparat"},
    "typ": {"typ", "type"},
    "charge": {"charge"},
    "verfall": {"verfall", "verfallsdatum"},
    "datum": {"datum", "date"},
    "anzahl": {"anzahl", "menge", "quantity"},
    "empfaenger": {"empfaenger", "empfänger"},
}


def _sanitize_filename_part(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", value or "")
    return safe.strip("._") or "datei"


def _build_attachment_target_path(
    base_dir: Path,
    bewegung_id: int,
    depot_name: str | None,
    original_name: str,
) -> Path:
    now = datetime.now()
    folder = base_dir / _sanitize_filename_part(depot_name or "depot") / now.strftime("%Y") / now.strftime("%m")
    folder.mkdir(parents=True, exist_ok=True)
    timestamp = now.strftime("%Y%m%d_%H%M%S")
    stem = _sanitize_filename_part(Path(original_name).stem or "anhang")
    filename = f"bewegung_{bewegung_id}_{timestamp}_{stem}.pdf"
    return folder / filename


def _persist_pdf_upload(upload_file: UploadFile, target: Path) -> int:
    source = upload_file.file
    source.seek(0)
    total_written = 0
    with target.open("wb") as out:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            total_written += len(chunk)
            if total_written > MAX_ATTACHMENT_SIZE_BYTES:
                out.close()
                target.unlink(missing_ok=True)
                raise ValueError("PDF-Datei ist zu gross (max. 10 MB).")
            out.write(chunk)
    return total_written


def _normalize_import_column_name(name: str) -> str:
    lowered = (name or "").strip().lower()
    lowered = lowered.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    lowered = re.sub(r"[^a-z0-9]+", "", lowered)
    return lowered


def _parse_id_list_csv(value: str) -> list[int]:
    if not value.strip():
        return []
    result: list[int] = []
    for part in value.split(","):
        item = part.strip()
        if not item:
            continue
        result.append(int(item))
    return result


def _add_months(iso_date: date, months: int) -> date:
    month_index = (iso_date.year * 12 + iso_date.month - 1) + months
    year = month_index // 12
    month = (month_index % 12) + 1
    return date(year, month, 1)


def _normalize_verfall_thresholds(
    critical_days: int = 30,
    warning_days: int = 90,
    attention_days: int = 180,
) -> tuple[int, int, int]:
    critical = max(1, min(int(critical_days), 3650))
    warning = max(critical + 1, min(int(warning_days), 3650))
    attention = max(warning + 1, min(int(attention_days), 3650))
    return critical, warning, attention


def _verfall_category(
    tage_bis_verfall: int,
    critical_days: int,
    warning_days: int,
    attention_days: int,
) -> str:
    if tage_bis_verfall <= critical_days:
        return "kritisch"
    if tage_bis_verfall <= warning_days:
        return "warnung"
    if tage_bis_verfall <= attention_days:
        return "achtung"
    return "ok"


def _enrich_verfall_rows(
    rows: list[dict],
    critical_days: int,
    warning_days: int,
    attention_days: int,
) -> list[dict]:
    enriched: list[dict] = []
    for row in rows:
        item = dict(row)
        try:
            tage = int(item.get("tage_bis_verfall"))
        except (TypeError, ValueError):
            tage = 99999
        item["tage_bis_verfall"] = tage
        item["kategorie"] = _verfall_category(tage, critical_days, warning_days, attention_days)
        enriched.append(item)
    return enriched


def _ensure_import_dependencies() -> None:
    if pd is None:
        raise ValueError("Import-Feature benoetigt pandas. Bitte Abhaengigkeit installieren.")
    if Workbook is None or DataValidation is None:
        raise ValueError("Import-Feature benoetigt openpyxl. Bitte Abhaengigkeit installieren.")


def _load_import_dataframe(upload_file: UploadFile) -> tuple[pd.DataFrame, int]:
    _ensure_import_dependencies()
    filename = (upload_file.filename or "").lower()
    upload_file.file.seek(0)
    if filename.endswith(".csv"):
        df = pd.read_csv(upload_file.file)
        return df, 2
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        upload_file.file.seek(0)
        df = pd.read_excel(upload_file.file, header=2)
        normalized_columns = {_normalize_import_column_name(str(col)) for col in df.columns}
        normalized_required = {
            _normalize_import_column_name(alias)
            for alias in IMPORT_COLUMN_ALIASES["depot"]
        }
        if not normalized_required.intersection(normalized_columns):
            upload_file.file.seek(0)
            df = pd.read_excel(upload_file.file, header=0)
            return df, 2
        return df, 4
    raise ValueError("Nur CSV-, XLSX- oder XLS-Dateien sind erlaubt.")


def _map_import_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename_map: dict[str, str] = {}
    for source_col in df.columns:
        normalized = _normalize_import_column_name(str(source_col))
        for target_col, aliases in IMPORT_COLUMN_ALIASES.items():
            normalized_aliases = {_normalize_import_column_name(alias) for alias in aliases}
            if normalized in normalized_aliases:
                rename_map[source_col] = target_col
                break
    mapped = df.rename(columns=rename_map).copy()
    missing = [col for col in IMPORT_REQUIRED_COLUMNS if col not in mapped.columns]
    if missing:
        raise ValueError(f"Fehlende Spalten: {', '.join(missing)}")
    return mapped


def _analyze_import_dataframe(
    repository: SqliteRepository,
    mapped_df: pd.DataFrame,
    row_offset: int,
) -> tuple[list[dict], list[dict], list[str]]:
    working = mapped_df.dropna(how="all").copy()
    if "depot" in working.columns:
        working = working[working["depot"].notna()]
        working = working[working["depot"].astype(str).str.strip() != ""]

    depots = repository.list_depots(limit=10000, offset=0, q="")
    depots_by_name = {str(row["name"]).strip().lower(): int(row["id"]) for row in depots}
    praeparate = repository.list_praeparate(limit=10000, offset=0, q="")
    prae_by_name = {str(row["name"]).strip().lower(): int(row["id"]) for row in praeparate}

    parsed_rows: list[dict] = []
    errors: list[str] = []

    for idx, row in working.iterrows():
        display_row = int(idx) + row_offset
        depot_name = str(row.get("depot", "")).strip()
        prae_name = str(row.get("praeparat", "")).strip()
        typ = str(row.get("typ", "")).strip()
        charge = str(row.get("charge", "")).strip()
        empfaenger_raw = row.get("empfaenger", None)
        empfaenger = str(empfaenger_raw).strip() if pd.notna(empfaenger_raw) else None

        depot_id = depots_by_name.get(depot_name.lower())
        if depot_id is None:
            errors.append(f"Zeile {display_row}: Depot '{depot_name}' nicht gefunden")
            continue
        praeparat_id = prae_by_name.get(prae_name.lower())
        if praeparat_id is None:
            errors.append(f"Zeile {display_row}: Praeparat '{prae_name}' nicht gefunden")
            continue
        if typ not in ALLOWED_IMPORT_TYPES:
            errors.append(f"Zeile {display_row}: Ungueltiger Typ '{typ}'")
            continue
        if not charge:
            errors.append(f"Zeile {display_row}: Charge darf nicht leer sein")
            continue
        try:
            anzahl = int(row.get("anzahl"))
            if anzahl <= 0:
                raise ValueError()
        except (ValueError, TypeError):
            errors.append(f"Zeile {display_row}: Anzahl muss > 0 sein")
            continue
        datum_iso = _parse_import_date(row.get("datum"))
        if datum_iso is None:
            errors.append(f"Zeile {display_row}: Datum ist ungueltig")
            continue
        verfall_iso = _parse_import_date(row.get("verfall"))
        if verfall_iso is None:
            errors.append(f"Zeile {display_row}: Verfall ist ungueltig")
            continue
        parsed_rows.append(
            {
                "display_row": display_row,
                "depot_id": depot_id,
                "praeparat_id": praeparat_id,
                "typ": typ,
                "charge": charge,
                "verfall": verfall_iso,
                "datum": datum_iso,
                "anzahl": anzahl,
                "empfaenger": empfaenger or None,
                "depot": depot_name,
                "praeparat": prae_name,
            }
        )

    preview_rows = working.head(10).fillna("").to_dict(orient="records")
    for row in preview_rows:
        for key, value in row.items():
            row[key] = str(value)
    return parsed_rows, preview_rows, errors


def _parse_import_date(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    value_text = str(value).strip()
    if not value_text:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value_text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    parsed = pd.to_datetime(value_text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.strftime("%Y-%m-%d")


def _build_import_template(repository: SqliteRepository) -> bytes:
    _ensure_import_dependencies()
    wb = Workbook()
    ws = wb.active
    ws.title = "Bewegungen"
    depots = [row["name"] for row in repository.list_depots(limit=10000, offset=0, q="")]
    praeparate = [row["name"] for row in repository.list_praeparate(limit=10000, offset=0, q="")]
    typen = ["Zugang", "Abgang", "Vernichtung"]

    validation = wb.create_sheet("Validierung")
    for idx, name in enumerate(depots, start=1):
        validation.cell(row=idx, column=1, value=name)
    for idx, name in enumerate(praeparate, start=1):
        validation.cell(row=idx, column=2, value=name)
    for idx, name in enumerate(typen, start=1):
        validation.cell(row=idx, column=3, value=name)
    validation.sheet_state = "hidden"

    ws.merge_cells("A1:H1")
    ws["A1"] = "Depot in B2 waehlen, danach Daten ab Zeile 4 eintragen."
    ws["A2"] = "Depot:"
    ws["B2"] = depots[0] if depots else ""
    headers = ["Depot", "Praeparat", "Typ", "Charge", "Verfall", "Datum", "Anzahl", "Empfaenger"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=header)
    for row_idx in range(4, 104):
        ws.cell(row=row_idx, column=1, value=f'=IF(B{row_idx}<>"",$B$2,"")')

    if depots:
        dv_depot = DataValidation(type="list", formula1=f"Validierung!$A$1:$A${len(depots)}", allow_blank=False)
        ws.add_data_validation(dv_depot)
        dv_depot.add("B2")
    if praeparate:
        dv_praep = DataValidation(type="list", formula1=f"Validierung!$B$1:$B${len(praeparate)}", allow_blank=False)
        ws.add_data_validation(dv_praep)
        dv_praep.add("B4:B103")
    dv_typ = DataValidation(type="list", formula1=f"Validierung!$C$1:$C${len(typen)}", allow_blank=False)
    ws.add_data_validation(dv_typ)
    dv_typ.add("C4:C103")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _csv_response(filename: str, headers: list[str], rows: list[list[object]]) -> StreamingResponse:
    text_io = io.StringIO()
    writer = csv.writer(text_io)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    buffer = BytesIO(text_io.getvalue().encode("utf-8"))
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pdf_response(filename: str, title: str, headers: list[str], rows: list[list[object]], kpis: dict) -> StreamingResponse:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PDF-Export erfordert reportlab.") from exc

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Heading1"]), Spacer(1, 8)]
    if kpis:
        for key, value in kpis.items():
            elements.append(Paragraph(f"<b>{key}:</b> {value}", styles["Normal"]))
        elements.append(Spacer(1, 8))
    table_data = [headers] + rows
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#007aff")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pptx_response(filename: str, title: str, headers: list[str], rows: list[list[object]], kpis: dict) -> StreamingResponse:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="PPT-Export erfordert python-pptx.") from exc

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_box = title_slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.size = Pt(24)
    title_tf.paragraphs[0].font.bold = True

    y = 1.0
    if kpis:
        kpi_box = title_slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(8.8), Inches(1.4))
        kpi_tf = kpi_box.text_frame
        first = True
        for key, value in kpis.items():
            p = kpi_tf.paragraphs[0] if first else kpi_tf.add_paragraph()
            first = False
            p.text = f"{key}: {value}"
            p.font.size = Pt(14)
        y += 1.5

    max_rows = min(len(rows), 18)
    cols = len(headers)
    table_shape = title_slide.shapes.add_table(cols=cols, rows=max_rows + 1, left=Inches(0.4), top=Inches(y), width=Inches(9.0), height=Inches(5.0))
    table = table_shape.table
    for col_idx, header in enumerate(headers):
        table.cell(0, col_idx).text = str(header)
    for r in range(max_rows):
        for c in range(cols):
            table.cell(r + 1, c).text = str(rows[r][c])

    with tempfile.NamedTemporaryFile(suffix=".pptx") as tmp:
        prs.save(tmp.name)
        tmp.seek(0)
        payload = tmp.read()
    out = BytesIO(payload)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


ALLOWED_USER_ROLES = {"Admin", "User"}
PERMISSION_DEFINITIONS = [
    {"key": "masterdata_read", "label": "Stammdaten lesen"},
    {"key": "masterdata_write", "label": "Stammdaten bearbeiten"},
    {"key": "movements_read", "label": "Bewegungen lesen"},
    {"key": "movements_write", "label": "Bewegungen erfassen"},
    {"key": "settings_read", "label": "Grundeinstellungen lesen"},
    {"key": "settings_write", "label": "Grundeinstellungen bearbeiten"},
    {"key": "import_use", "label": "Import nutzen"},
    {"key": "email_use", "label": "E-Mail nutzen"},
    {"key": "reports_view", "label": "Auswertungen sehen"},
    {"key": "audit_view", "label": "Audit sehen"},
    {"key": "users_manage", "label": "Benutzer verwalten"},
    {"key": "backup_manage", "label": "Backup verwalten"},
]
ALL_PERMISSION_KEYS = {item["key"] for item in PERMISSION_DEFINITIONS}
DEFAULT_USER_PERMISSIONS = {
    "masterdata_read",
    "movements_read",
    "movements_write",
    "settings_read",
    "import_use",
    "email_use",
    "reports_view",
}
PERMISSION_TEMPLATES = [
    {"key": "readonly", "label": "ReadOnly", "permissions": ["masterdata_read", "movements_read", "settings_read", "reports_view"]},
    {"key": "disponent", "label": "Disponent", "permissions": ["masterdata_read", "movements_read", "movements_write", "settings_read", "import_use", "email_use", "reports_view"]},
    {"key": "reporting", "label": "Reporting", "permissions": ["masterdata_read", "movements_read", "reports_view"]},
    {"key": "ops_admin", "label": "Ops Admin", "permissions": ["masterdata_read", "masterdata_write", "movements_read", "movements_write", "settings_read", "settings_write", "import_use", "email_use", "reports_view", "audit_view", "backup_manage"]},
]


def _parse_permission_list(raw_permissions: object) -> set[str]:
    if raw_permissions in (None, ""):
        return set()
    if isinstance(raw_permissions, str):
        try:
            parsed = json.loads(raw_permissions)
        except json.JSONDecodeError:
            return set()
    elif isinstance(raw_permissions, (list, tuple, set)):
        parsed = list(raw_permissions)
    else:
        return set()
    result: set[str] = set()
    for item in parsed:
        key = str(item or "").strip()
        if key in ALL_PERMISSION_KEYS:
            result.add(key)
    return result


def _permissions_for_role(role: str | None, raw_permissions: object) -> set[str]:
    if role == "Admin":
        return set(ALL_PERMISSION_KEYS)
    parsed = _parse_permission_list(raw_permissions)
    return parsed or set(DEFAULT_USER_PERMISSIONS)


def _permissions_json_for_storage(role: str, requested: list[str] | None) -> str:
    allowed = _permissions_for_role(role, requested)
    return json.dumps(sorted(allowed), ensure_ascii=True)


def _normalize_user_row(row: tuple) -> dict:
    effective_permissions = sorted(_permissions_for_role(str(row[2]), row[10] if len(row) > 10 else None))
    return {
        "id": int(row[0]),
        "username": str(row[1]),
        "role": str(row[2]),
        "email": row[3],
        "created_at": row[4],
        "last_login": row[5],
        "is_active": bool(row[6]),
        "failed_attempts": int(row[7] or 0),
        "locked_until": row[8],
        "is_default_password": bool(row[9]),
        "permissions": effective_permissions,
    }


def _get_user_flags(security: SecurityManager, username: str) -> dict[str, bool]:
    row = security.cur.execute(
        "SELECT is_default_password, role, permissions FROM users WHERE username = ?",
        ((username or "").strip(),),
    ).fetchone()
    if not row:
        return {"requires_password_change": False, "permissions": sorted(DEFAULT_USER_PERMISSIONS)}  # nosec B105: boolean flag, not a password
    return {
        "requires_password_change": bool(row[0]),
        "permissions": sorted(_permissions_for_role(str(row[1]), row[2])),
    }


def _get_avatar_path_for_user(security: SecurityManager, username: str) -> Path | None:
    row = security.cur.execute(
        "SELECT avatar_path FROM users WHERE username = ?",
        ((username or "").strip(),),
    ).fetchone()
    if not row or not row[0]:
        return None
    avatar_path = Path(str(row[0])).expanduser()
    if not avatar_path.exists() or not avatar_path.is_file():
        return None
    return avatar_path


def _safe_backup_label(value: str) -> str:
    return re.sub(r"[^a-z0-9_-]+", "_", (value or "").strip().lower()) or "manual"


def _create_backup_snapshot(source_db_path: Path, backups_dir: Path, label: str) -> Path:
    backups_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"ndhub_{_safe_backup_label(label)}_{timestamp}.db"
    target_path = backups_dir / filename
    with sqlite3.connect(str(source_db_path)) as source_conn:
        with sqlite3.connect(str(target_path)) as target_conn:
            source_conn.backup(target_conn)
    return target_path


def _list_backup_files(backups_dir: Path, limit: int = 200) -> list[dict]:
    if not backups_dir.exists():
        return []
    rows: list[dict] = []
    for path in sorted(backups_dir.glob("ndhub_*.db"), key=lambda item: item.stat().st_mtime, reverse=True):
        stat = path.stat()
        rows.append(
            {
                "filename": path.name,
                "size_bytes": int(stat.st_size),
                "created_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
            }
        )
        if len(rows) >= max(1, min(int(limit), 1000)):
            break
    return rows


def _run_auto_backup_if_due(source_db_path: Path, backups_dir: Path, interval_hours: int) -> Path | None:
    safe_interval_hours = max(1, min(int(interval_hours), 24 * 30))
    now_ts = datetime.now(timezone.utc).timestamp()
    latest_auto_mtime = None
    for path in backups_dir.glob("ndhub_auto_*.db"):
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        if latest_auto_mtime is None or mtime > latest_auto_mtime:
            latest_auto_mtime = mtime
    if latest_auto_mtime is not None:
        elapsed_hours = (now_ts - latest_auto_mtime) / 3600
        if elapsed_hours < safe_interval_hours:
            return None
    return _create_backup_snapshot(source_db_path, backups_dir, "auto")


def create_app(db_path: str | None = None) -> FastAPI:
    """Application factory for runtime and tests."""
    database_path = db_path or resolve_db_path()
    source_db_path = Path(database_path).resolve()
    backups_dir = source_db_path.parent / "backups"
    auto_backup_hours = int((os.environ.get("ND_HUB_AUTO_BACKUP_HOURS", "24") or "24").strip() or "24")
    repository = SqliteRepository(database_path)
    security = SecurityManager(database_path)
    token_store = TokenStore()

    def _user_permissions(username: str, role: str | None) -> set[str]:
        if role == "Admin":
            return set(ALL_PERMISSION_KEYS)
        row = security.cur.execute(
            "SELECT permissions FROM users WHERE username = ?",
            ((username or "").strip(),),
        ).fetchone()
        raw_permissions = row[0] if row else None
        return _permissions_for_role(role, raw_permissions)

    def require_permission(permission_key: str):
        def _dependency(session: SessionInfo = Depends(get_current_session)) -> SessionInfo:
            if permission_key not in ALL_PERMISSION_KEYS:
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail=f"Unbekannte Berechtigung: {permission_key}",
                )
            if session.role == "Admin":
                return session
            permissions = _user_permissions(session.username, session.role)
            if permission_key not in permissions:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Berechtigung fehlt.",
                )
            return session

        return _dependency

    @asynccontextmanager
    async def _lifespan(_app: FastAPI):
        try:
            try:
                created = _run_auto_backup_if_due(source_db_path, backups_dir, auto_backup_hours)
                if created is not None:
                    repository.log_audit(
                        username="system",
                        action="create",
                        resource_type="backup",
                        details={"operation": "auto", "filename": created.name},
                    )
            except Exception:
                # Auto backup should never block startup.
                logger.warning("Auto-backup during startup failed", exc_info=True)
            yield
        finally:
            security.close()

    app = FastAPI(title="ND-Hub Backend", version="0.1.1", lifespan=_lifespan)
    app.state.repository = repository
    app.state.security = security
    app.state.token_store = token_store
    app.state.database_path = database_path
    app.state.web_dir = Path(__file__).parent / "web"
    app.state.attachments_dir = Path(database_path).resolve().parent / "attachments"
    app.state.backups_dir = backups_dir
    app.state.auto_backup_hours = auto_backup_hours

    # ---- Security-Middleware (Issue #68 — Port aus Web-Backend) ----
    # 1) TrustedHost: blockiert Host-Header-Injection
    _allowed_hosts_raw = os.environ.get("ND_HUB_ALLOWED_HOSTS", "*").strip()
    if _allowed_hosts_raw == "*":
        _allowed_hosts = ["*"]
    else:
        _allowed_hosts = [h.strip() for h in _allowed_hosts_raw.split(",") if h.strip()]
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=_allowed_hosts)

    # 2) CORS: restriktiver Default (leere allow_origins)
    _cors_raw = os.environ.get("ND_HUB_CORS_ORIGINS", "").strip()
    _cors_origins = [o.strip() for o in _cors_raw.split(",") if o.strip()] if _cors_raw else []
    app.add_middleware(
        CORSMiddleware,
        allow_origins=_cors_origins,
        allow_credentials=True,
        allow_methods=["GET", "POST", "PUT", "DELETE", "PATCH", "OPTIONS"],
        allow_headers=["Authorization", "Content-Type", "Accept"],
    )

    # 3) Security-Header auf jede Response
    @app.middleware("http")
    async def _security_headers_middleware(request: Request, call_next):
        response = await call_next(request)
        response.headers.setdefault("X-Content-Type-Options", "nosniff")
        response.headers.setdefault("X-Frame-Options", "DENY")
        response.headers.setdefault("Referrer-Policy", "no-referrer")
        if request.url.scheme == "https":
            response.headers.setdefault(
                "Strict-Transport-Security",
                "max-age=63072000; includeSubDomains",
            )
        return response
    # ---- /Security-Middleware ----

    # ---- Login-Rate-Limit (Issue #68 — Port aus Web-Backend) ----
    import threading as _threading

    app.state.login_lockout_state: dict[str, tuple[int, float]] = {}
    app.state.login_lockout_lock = _threading.Lock()
    app.state.login_max_fails = int(os.environ.get("ND_HUB_LOGIN_MAX_FAILS", "10"))
    app.state.login_lockout_seconds = int(os.environ.get("ND_HUB_LOGIN_LOCKOUT_SECONDS", "900"))
    # ---- /Login-Rate-Limit ----

    app.mount(
        "/web",
        StaticFiles(directory=app.state.web_dir, html=False),
        name="web",
    )

    @app.get("/")
    def web_index() -> FileResponse:
        return FileResponse(app.state.web_dir / "index.html")

    @app.get("/health")
    def health() -> dict[str, str]:
        return {"status": "ok"}

    @app.post("/auth/login", response_model=LoginResponse)
    def login(payload: LoginRequest, request: Request) -> LoginResponse:
        # ---- Login-Lockout (Issue #68) ----
        username = payload.username.strip()
        with app.state.login_lockout_lock:
            _state = app.state.login_lockout_state
            import time as _time
            now = _time.time()
            if username in _state:
                fails, lockout_until = _state[username]
                if lockout_until > now:
                    raise HTTPException(
                        status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                        detail="Account temporär gesperrt. Bitte später erneut versuchen.",
                    )

        ok, _message = security.authenticate(username, payload.password)
        if not ok:
            # Fehlversuch registrieren
            with app.state.login_lockout_lock:
                import time as _time
                now = _time.time()
                if username in _state:
                    fails, _lockout = _state[username]
                else:
                    fails = 0
                fails += 1
                if fails >= app.state.login_max_fails:
                    _state[username] = (fails, now + app.state.login_lockout_seconds)
                else:
                    _state[username] = (fails, 0.0)
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Benutzername oder Passwort ungueltig.",
            )

        # Erfolg → Lockout-Counter zurücksetzen
        with app.state.login_lockout_lock:
            _state.pop(username, None)

        user_flags = _get_user_flags(security, username)
        token = token_store.issue(
            username=username,
            role=security.get_current_role(),
        )
        return LoginResponse(
            token=token,
            username=username,
            role=security.get_current_role(),
            requires_password_change=user_flags["requires_password_change"],
            permissions=user_flags["permissions"],
        )

    @app.get("/auth/me")
    def auth_me(session: SessionInfo = Depends(get_current_session)) -> dict[str, object]:
        user_flags = _get_user_flags(security, session.username)
        avatar_path = _get_avatar_path_for_user(security, session.username)
        return {
            "username": session.username,
            "role": session.role,
            "expires_at": session.expires_at.isoformat(),
            "requires_password_change": user_flags["requires_password_change"],
            "permissions": user_flags["permissions"],
            "avatar_available": avatar_path is not None,
        }

    @app.get("/permissions/catalog")
    def permissions_catalog(
        session: SessionInfo = Depends(get_current_session),
    ) -> dict:
        _ = session
        return {
            "rows": PERMISSION_DEFINITIONS,
            "default_user_permissions": sorted(DEFAULT_USER_PERMISSIONS),
            "templates": PERMISSION_TEMPLATES,
        }

    @app.get("/auth/activity")
    def auth_activity(
        limit: int = 50,
        session: SessionInfo = Depends(get_current_session),
    ) -> list[dict]:
        safe_limit = max(1, min(int(limit), 200))
        user_row = security.cur.execute(
            "SELECT id FROM users WHERE username = ?",
            (session.username,),
        ).fetchone()
        if not user_row:
            return []
        rows = security.get_activity_log(user_id=int(user_row[0]), limit=safe_limit)
        result: list[dict] = []
        for row in rows:
            result.append(
                {
                    "id": int(row[0]),
                    "username": str(row[1]),
                    "action": str(row[2]),
                    "details": str(row[3] or ""),
                    "timestamp": str(row[4]),
                }
            )
        return result

    @app.get("/auth/avatar")
    def auth_avatar(
        session: SessionInfo = Depends(get_current_session),
    ) -> FileResponse:
        avatar_path = _get_avatar_path_for_user(security, session.username)
        if avatar_path is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Kein Profilbild vorhanden.")
        return FileResponse(path=avatar_path)

    @app.post("/auth/avatar")
    def upload_auth_avatar(
        file: UploadFile = File(...),
        session: SessionInfo = Depends(get_current_session),
    ) -> dict[str, str]:
        filename = (file.filename or "").strip()
        if not filename:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dateiname fehlt.")
        suffix = Path(filename).suffix.lower()
        if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Nur PNG/JPG/JPEG/WEBP/BMP sind erlaubt.",
            )
        user_row = security.cur.execute(
            "SELECT id FROM users WHERE username = ?",
            (session.username,),
        ).fetchone()
        if not user_row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp_path = Path(tmp.name)
            file.file.seek(0)
            shutil.copyfileobj(file.file, tmp)
        try:
            security.current_user = session.username
            security.current_role = session.role
            ok, message = security.set_user_avatar(int(user_row[0]), str(tmp_path))
        finally:
            tmp_path.unlink(missing_ok=True)
        if not ok:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=message)
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="user",
            resource_id=int(user_row[0]),
            details={"change": "avatar_upload"},
        )
        return {"status": "avatar_saved"}

    @app.delete("/auth/avatar")
    def clear_auth_avatar(
        session: SessionInfo = Depends(get_current_session),
    ) -> dict[str, str]:
        user_row = security.cur.execute(
            "SELECT id FROM users WHERE username = ?",
            (session.username,),
        ).fetchone()
        if not user_row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        security.current_user = session.username
        security.current_role = session.role
        ok, message = security.clear_user_avatar(int(user_row[0]))
        if not ok:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=message)
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="user",
            resource_id=int(user_row[0]),
            details={"change": "avatar_clear"},
        )
        return {"status": "avatar_cleared"}

    @app.post("/auth/logout")
    def auth_logout(
        session: SessionInfo = Depends(get_current_session),
        credentials: HTTPAuthorizationCredentials | None = Depends(bearer_scheme),
    ) -> dict[str, str]:
        _ = session
        if credentials and credentials.credentials:
            token_store.revoke(credentials.credentials)
        return {"status": "logged_out"}

    @app.post("/auth/change-password")
    def auth_change_password(
        payload: PasswordChangeRequest,
        session: SessionInfo = Depends(get_current_session),
    ) -> dict[str, str]:
        user_row = security.cur.execute(
            "SELECT id FROM users WHERE username = ?",
            (session.username,),
        ).fetchone()
        if not user_row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        ok, message = security.change_password(int(user_row[0]), payload.old_password, payload.new_password)
        if not ok:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=message)
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="user",
            resource_id=int(user_row[0]),
            details={"change": "password_self_service"},
        )
        return {"status": "password_changed"}

    @app.get("/users")
    def list_users(
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> list[dict]:
        _ = session
        return [_normalize_user_row(row) for row in security.list_users()]

    @app.get("/users/{user_id}/activity")
    def list_user_activity(
        user_id: int,
        limit: int = 100,
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> list[dict]:
        _ = session
        safe_limit = max(1, min(int(limit), 500))
        rows = security.get_activity_log(user_id=int(user_id), limit=safe_limit)
        result: list[dict] = []
        for row in rows:
            result.append(
                {
                    "id": int(row[0]),
                    "username": str(row[1]),
                    "action": str(row[2]),
                    "details": str(row[3] or ""),
                    "timestamp": str(row[4]),
                }
            )
        return result

    @app.post("/users", status_code=status.HTTP_201_CREATED)
    def create_user(
        payload: UserCreateRequest,
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> dict:
        _ = session
        safe_username = payload.username.strip()
        if not safe_username:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Benutzername darf nicht leer sein.")
        safe_role = payload.role.strip()
        if safe_role not in ALLOWED_USER_ROLES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Rolle.")
        exists = security.cur.execute(
            "SELECT id FROM users WHERE username = ?",
            (safe_username,),
        ).fetchone()
        if exists:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Benutzername existiert bereits.")
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        password_hash = SecurityManager.hash_password(payload.password)
        permissions_json = _permissions_json_for_storage(safe_role, payload.permissions)
        security.cur.execute(
            """
            INSERT INTO users (username, password_hash, role, email, created_at, is_active, is_default_password, permissions)
            VALUES (?, ?, ?, ?, ?, ?, 0, ?)
            """,
            (
                safe_username,
                password_hash,
                safe_role,
                (payload.email or "").strip() or None,
                now,
                1 if payload.is_active else 0,
                permissions_json,
            ),
        )
        security.conn.commit()
        user_id = int(security.cur.lastrowid)
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="user",
            resource_id=user_id,
            details={"username": safe_username, "role": safe_role},
        )
        row = security.cur.execute(
            """
            SELECT id, username, role, email, created_at, last_login, is_active, failed_attempts, locked_until, is_default_password, permissions
            FROM users
            WHERE id = ?
            """,
            (user_id,),
        ).fetchone()
        return _normalize_user_row(row)

    @app.put("/users/{user_id}")
    def update_user(
        user_id: int,
        payload: UserUpdateRequest,
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> dict:
        _ = session
        row = security.cur.execute(
            "SELECT id, username, role, is_active FROM users WHERE id = ?",
            (int(user_id),),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        updates: list[str] = []
        params: list[object] = []
        if payload.username is not None:
            safe_username = payload.username.strip()
            if not safe_username:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Benutzername darf nicht leer sein.")
            duplicate = security.cur.execute(
                "SELECT id FROM users WHERE username = ? AND id != ?",
                (safe_username, int(user_id)),
            ).fetchone()
            if duplicate:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Benutzername existiert bereits.")
            updates.append("username = ?")
            params.append(safe_username)
        if payload.role is not None:
            safe_role = payload.role.strip()
            if safe_role not in ALLOWED_USER_ROLES:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Rolle.")
            updates.append("role = ?")
            params.append(safe_role)
        if payload.email is not None:
            updates.append("email = ?")
            params.append((payload.email or "").strip() or None)
        if payload.permissions is not None:
            effective_role = payload.role.strip() if payload.role else str(row[2])
            updates.append("permissions = ?")
            params.append(_permissions_json_for_storage(effective_role, payload.permissions))
        if payload.is_active is not None:
            current_role = str(row[2])
            if current_role == "Admin" and payload.is_active is False:
                admin_count = security.cur.execute(
                    "SELECT COUNT(*) FROM users WHERE role = 'Admin' AND is_active = 1",
                ).fetchone()[0]
                if int(admin_count) <= 1:
                    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Letzter aktiver Admin darf nicht deaktiviert werden.")
            updates.append("is_active = ?")
            params.append(1 if payload.is_active else 0)
        if not updates:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Keine Änderungen angegeben.")
        params.append(int(user_id))
        # SET clauses are hardcoded "<column> = ?" strings from an explicit
        # allow-list; only the values are parameterized and user-controlled.
        security.cur.execute(
            f"UPDATE users SET {', '.join(updates)} WHERE id = ?",  # nosec B608
            tuple(params),
        )
        security.conn.commit()
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="user",
            resource_id=int(user_id),
            details={"updated_fields": updates},
        )
        refreshed = security.cur.execute(
            """
            SELECT id, username, role, email, created_at, last_login, is_active, failed_attempts, locked_until, is_default_password, permissions
            FROM users
            WHERE id = ?
            """,
            (int(user_id),),
        ).fetchone()
        return _normalize_user_row(refreshed)

    @app.delete("/users/{user_id}")
    def delete_user(
        user_id: int,
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> dict[str, int | str]:
        _ = session
        row = security.cur.execute(
            "SELECT id, username, role, is_active FROM users WHERE id = ?",
            (int(user_id),),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        if str(row[1]) == session.username:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Sie können sich nicht selbst löschen.")
        if str(row[2]) == "Admin" and int(row[3]) == 1:
            admin_count = security.cur.execute(
                "SELECT COUNT(*) FROM users WHERE role = 'Admin' AND is_active = 1",
            ).fetchone()[0]
            if int(admin_count) <= 1:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Letzter aktiver Admin darf nicht gelöscht werden.")
        security.cur.execute("DELETE FROM users WHERE id = ?", (int(user_id),))
        security.conn.commit()
        repository.log_audit(
            username=session.username,
            action="delete",
            resource_type="user",
            resource_id=int(user_id),
            details={"username": row[1]},
        )
        return {"id": int(user_id), "status": "deleted"}

    @app.post("/users/{user_id}/reset-password")
    def reset_user_password(
        user_id: int,
        payload: UserPasswordResetRequest,
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> dict[str, int | str]:
        _ = session
        row = security.cur.execute(
            "SELECT id, username FROM users WHERE id = ?",
            (int(user_id),),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        new_hash = SecurityManager.hash_password(payload.new_password)
        security.cur.execute(
            """
            UPDATE users
            SET password_hash = ?, failed_attempts = 0, locked_until = NULL, is_default_password = 1
            WHERE id = ?
            """,
            (new_hash, int(user_id)),
        )
        security.conn.commit()
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="user",
            resource_id=int(user_id),
            details={"change": "password_reset"},
        )
        return {"id": int(user_id), "status": "password_reset"}

    @app.post("/users/{user_id}/unlock")
    def unlock_user(
        user_id: int,
        session: SessionInfo = Depends(require_permission("users_manage")),
    ) -> dict[str, int | str]:
        _ = session
        row = security.cur.execute(
            "SELECT id FROM users WHERE id = ?",
            (int(user_id),),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Benutzer nicht gefunden.")
        security.cur.execute(
            "UPDATE users SET failed_attempts = 0, locked_until = NULL WHERE id = ?",
            (int(user_id),),
        )
        security.conn.commit()
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="user",
            resource_id=int(user_id),
            details={"change": "unlock"},
        )
        return {"id": int(user_id), "status": "unlocked"}

    @app.get("/admin/backup/list")
    def list_backups(
        limit: int = 200,
        session: SessionInfo = Depends(require_permission("backup_manage")),
    ) -> dict:
        _ = session
        rows = _list_backup_files(backups_dir, limit=limit)
        return {"rows": rows, "auto_backup_hours": auto_backup_hours}

    @app.post("/admin/backup/create")
    def create_backup(
        session: SessionInfo = Depends(require_permission("backup_manage")),
    ) -> dict:
        _ = session
        backup_path = _create_backup_snapshot(source_db_path, backups_dir, "manual")
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="backup",
            details={"operation": "manual", "filename": backup_path.name},
        )
        stat = backup_path.stat()
        return {
            "filename": backup_path.name,
            "size_bytes": int(stat.st_size),
            "created_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
        }

    @app.get("/admin/backup/download")
    def download_backup(
        session: SessionInfo = Depends(require_permission("backup_manage")),
    ) -> FileResponse:
        _ = session
        backup_path = _create_backup_snapshot(source_db_path, backups_dir, "manual")
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="backup",
            details={"operation": "manual_download", "filename": backup_path.name},
        )
        return FileResponse(
            path=backup_path,
            filename=backup_path.name,
            media_type="application/octet-stream",
        )

    @app.get("/admin/backup/download/{filename}")
    def download_backup_file(
        filename: str,
        session: SessionInfo = Depends(require_permission("backup_manage")),
    ) -> FileResponse:
        _ = session
        safe_name = Path(filename).name
        backup_path = backups_dir / safe_name
        if not backup_path.exists() or not backup_path.is_file():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Backup-Datei nicht gefunden.")
        return FileResponse(
            path=backup_path,
            filename=safe_name,
            media_type="application/octet-stream",
        )

    @app.post("/admin/backup/restore")
    def restore_backup(
        file: UploadFile = File(...),
        session: SessionInfo = Depends(require_permission("backup_manage")),
    ) -> dict[str, str]:
        nonlocal security
        nonlocal token_store
        _ = session
        filename = (file.filename or "").strip().lower()
        if not filename.endswith((".db", ".sqlite", ".sqlite3")):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Nur Backup-Dateien mit Endung .db/.sqlite/.sqlite3 sind erlaubt.",
            )
        db_target_path = Path(database_path).resolve()
        temp_restore_path = db_target_path.with_name(f"{db_target_path.stem}.restore_tmp{db_target_path.suffix}")
        file.file.seek(0)
        with temp_restore_path.open("wb") as temp_out:
            shutil.copyfileobj(file.file, temp_out)
        try:
            with sqlite3.connect(str(temp_restore_path)) as conn:
                required = {"users", "depots", "praeparate", "bewegungen"}
                rows = conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'",
                ).fetchall()
                present = {str(row[0]) for row in rows}
                if not required.issubset(present):
                    missing = ", ".join(sorted(required - present))
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"Backup ungueltig, fehlende Tabellen: {missing}",
                    )
        except HTTPException:
            temp_restore_path.unlink(missing_ok=True)
            raise
        except Exception as exc:
            temp_restore_path.unlink(missing_ok=True)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Backup-Datei konnte nicht validiert werden: {exc}",
            ) from exc

        pre_restore_path = _create_backup_snapshot(source_db_path, backups_dir, "pre_restore")
        try:
            security.close()
            shutil.move(str(temp_restore_path), str(db_target_path))
            security = SecurityManager(database_path)
            app.state.security = security
            token_store = TokenStore()
            app.state.token_store = token_store
        except Exception as exc:
            temp_restore_path.unlink(missing_ok=True)
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Restore fehlgeschlagen: {exc}",
            ) from exc

        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="backup",
            details={"operation": "restore", "source_file": filename, "pre_restore_file": pre_restore_path.name},
        )
        return {
            "status": "restored",
            "pre_restore_backup": pre_restore_path.name,
            "message": "Backup erfolgreich wiederhergestellt. Bitte neu anmelden.",
        }

    @app.get("/depots")
    def list_depots(
        q: str = "",
        limit: int = 100,
        offset: int = 0,
        session: SessionInfo = Depends(require_permission("masterdata_read")),
    ) -> list[dict]:
        _ = session
        return repository.list_depots(q=q, limit=limit, offset=offset)

    @app.post("/depots", status_code=status.HTTP_201_CREATED)
    def create_depot(
        payload: DepotUpsertRequest,
        session: SessionInfo = Depends(require_permission("masterdata_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            new_id = repository.create_depot(
                name=payload.name,
                adresse=payload.adresse,
                telefon=payload.telefon,
                email=payload.email,
            )
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="depot",
            resource_id=new_id,
            details={"name": payload.name},
        )
        return {"id": new_id, "status": "created"}

    @app.put("/depots/{depot_id}")
    def update_depot(
        depot_id: int,
        payload: DepotUpsertRequest,
        session: SessionInfo = Depends(require_permission("masterdata_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            changed = repository.update_depot(
                depot_id=depot_id,
                name=payload.name,
                adresse=payload.adresse,
                telefon=payload.telefon,
                email=payload.email,
            )
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        if not changed:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Depot nicht gefunden.")
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="depot",
            resource_id=depot_id,
            details={"name": payload.name},
        )
        return {"id": depot_id, "status": "updated"}

    @app.delete("/depots/{depot_id}")
    def delete_depot(
        depot_id: int,
        session: SessionInfo = Depends(require_permission("masterdata_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            changed = repository.delete_depot(depot_id)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        if not changed:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Depot nicht gefunden.")
        repository.log_audit(
            username=session.username,
            action="delete",
            resource_type="depot",
            resource_id=depot_id,
        )
        return {"id": depot_id, "status": "deleted"}

    @app.get("/praeparate")
    def list_praeparate(
        q: str = "",
        limit: int = 100,
        offset: int = 0,
        session: SessionInfo = Depends(require_permission("masterdata_read")),
    ) -> list[dict]:
        _ = session
        return repository.list_praeparate(q=q, limit=limit, offset=offset)

    @app.get("/depots/{depot_id}/praeparate")
    def list_praeparate_for_depot(
        depot_id: int,
        session: SessionInfo = Depends(require_permission("masterdata_read")),
    ) -> list[dict]:
        _ = session
        return repository.list_praeparate_for_depot(depot_id)

    @app.get("/depots/{depot_id}/zuordnungen")
    def list_depot_assignments(
        depot_id: int,
        session: SessionInfo = Depends(require_permission("settings_read")),
    ) -> list[dict]:
        _ = session
        return repository.list_depot_assignments(depot_id)

    @app.put("/depots/{depot_id}/zuordnungen")
    def update_depot_assignments(
        depot_id: int,
        payload: DepotAssignmentsUpdateRequest,
        session: SessionInfo = Depends(require_permission("settings_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            repository.set_depot_assignments(
                depot_id=depot_id,
                assignments=[item.model_dump() for item in payload.assignments],
            )
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="depot",
            resource_id=depot_id,
            details={"assignments_count": len(payload.assignments)},
        )
        return {"id": depot_id, "status": "assignments_updated"}

    @app.post("/praeparate", status_code=status.HTTP_201_CREATED)
    def create_praeparat(
        payload: PraeparatUpsertRequest,
        session: SessionInfo = Depends(require_permission("masterdata_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            new_id = repository.create_praeparat(name=payload.name)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="praeparat",
            resource_id=new_id,
            details={"name": payload.name},
        )
        return {"id": new_id, "status": "created"}

    @app.put("/praeparate/{praeparat_id}")
    def update_praeparat(
        praeparat_id: int,
        payload: PraeparatUpsertRequest,
        session: SessionInfo = Depends(require_permission("masterdata_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            changed = repository.update_praeparat(praeparat_id=praeparat_id, name=payload.name)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        if not changed:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Praeparat nicht gefunden.")
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="praeparat",
            resource_id=praeparat_id,
            details={"name": payload.name},
        )
        return {"id": praeparat_id, "status": "updated"}

    @app.delete("/praeparate/{praeparat_id}")
    def delete_praeparat(
        praeparat_id: int,
        session: SessionInfo = Depends(require_permission("masterdata_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            changed = repository.delete_praeparat(praeparat_id)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        if not changed:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Praeparat nicht gefunden.")
        repository.log_audit(
            username=session.username,
            action="delete",
            resource_type="praeparat",
            resource_id=praeparat_id,
        )
        return {"id": praeparat_id, "status": "deleted"}

    @app.get("/depots/{depot_id}/kontakte")
    def list_kontakte_for_depot(
        depot_id: int,
        session: SessionInfo = Depends(require_permission("settings_read")),
    ) -> list[dict]:
        _ = session
        return repository.list_kontakte(depot_id)

    @app.post("/depots/{depot_id}/kontakte", status_code=status.HTTP_201_CREATED)
    def create_kontakt_for_depot(
        depot_id: int,
        payload: KontaktUpsertRequest,
        session: SessionInfo = Depends(require_permission("settings_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            kontakt_id = repository.create_kontakt(
                depot_id=depot_id,
                name=payload.name,
                rolle=payload.rolle,
                telefon=payload.telefon,
                email=payload.email,
            )
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="kontakt",
            resource_id=kontakt_id,
            details={"depot_id": depot_id, "name": payload.name},
        )
        return {"id": kontakt_id, "status": "created"}

    @app.put("/kontakte/{kontakt_id}")
    def update_kontakt(
        kontakt_id: int,
        payload: KontaktUpsertRequest,
        session: SessionInfo = Depends(require_permission("settings_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            changed = repository.update_kontakt(
                kontakt_id=kontakt_id,
                name=payload.name,
                rolle=payload.rolle,
                telefon=payload.telefon,
                email=payload.email,
            )
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        if not changed:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Kontakt nicht gefunden.")
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="kontakt",
            resource_id=kontakt_id,
            details={"name": payload.name},
        )
        return {"id": kontakt_id, "status": "updated"}

    @app.delete("/kontakte/{kontakt_id}")
    def delete_kontakt(
        kontakt_id: int,
        session: SessionInfo = Depends(require_permission("settings_write")),
    ) -> dict[str, int | str]:
        _ = session
        changed = repository.delete_kontakt(kontakt_id)
        if not changed:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Kontakt nicht gefunden.")
        repository.log_audit(
            username=session.username,
            action="delete",
            resource_type="kontakt",
            resource_id=kontakt_id,
        )
        return {"id": kontakt_id, "status": "deleted"}

    @app.post("/emails/recipients-preview")
    def preview_email_recipients(
        payload: EmailRecipientPreviewRequest,
        session: SessionInfo = Depends(require_permission("email_use")),
    ) -> dict:
        _ = session
        if not payload.depot_ids:
            return {"count": 0, "recipients": [], "depot_names": []}
        recipients = repository.get_kontakte_by_depot_ids(payload.depot_ids)
        depot_names = sorted({str(row["depot_name"]) for row in recipients})
        return {"count": len(recipients), "recipients": recipients, "depot_names": depot_names}

    @app.post("/emails/drafts", status_code=status.HTTP_201_CREATED)
    def create_email_draft(
        payload: EmailDraftCreateRequest,
        session: SessionInfo = Depends(require_permission("email_use")),
    ) -> dict[str, int | str]:
        _ = session
        if not payload.depot_ids:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Bitte mindestens ein Depot auswaehlen.")
        recipients = repository.get_kontakte_by_depot_ids(payload.depot_ids)
        if not recipients:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Fuer die ausgewaehlten Depots sind keine Ansprechpartner mit E-Mail hinterlegt.",
            )
        emails = [str(row["email"]).strip() for row in recipients if str(row.get("email", "")).strip()]
        unique_emails = sorted(set(emails))
        depot_names = sorted({str(row["depot_name"]) for row in recipients})
        log_id = repository.add_email_verlauf(
            betreff=payload.betreff,
            nachricht=payload.nachricht or "",
            depot_names=", ".join(depot_names),
            emails="; ".join(unique_emails),
            anzahl=len(unique_emails),
        )
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="email",
            resource_id=log_id,
            details={"depot_ids": payload.depot_ids, "recipient_count": len(unique_emails)},
        )
        return {"id": log_id, "status": "draft_created", "recipient_count": len(unique_emails)}

    @app.get("/emails/history")
    def list_email_history(
        limit: int = 50,
        session: SessionInfo = Depends(require_permission("email_use")),
    ) -> list[dict]:
        _ = session
        return repository.get_email_verlauf(limit=limit)

    @app.get("/emails/history/{email_id}")
    def get_email_history_detail(
        email_id: int,
        session: SessionInfo = Depends(require_permission("email_use")),
    ) -> dict:
        _ = session
        details = repository.get_email_details(email_id)
        if details is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="E-Mail-Eintrag nicht gefunden.")
        return details

    @app.get("/bewegungen")
    def list_bewegungen(
        limit: int = 100,
        offset: int = 0,
        q: str = "",
        typ: str = "",
        session: SessionInfo = Depends(require_permission("movements_read")),
    ) -> list[dict]:
        _ = session
        return repository.list_bewegungen(limit=limit, offset=offset, q=q, typ=typ or None)

    @app.get("/dashboard/overview")
    def dashboard_overview(
        critical_days: int = 30,
        warning_days: int = 90,
        attention_days: int = 180,
        session: SessionInfo = Depends(require_permission("movements_read")),
    ) -> dict:
        _ = session
        safe_critical, safe_warning, safe_attention = _normalize_verfall_thresholds(
            critical_days=critical_days,
            warning_days=warning_days,
            attention_days=attention_days,
        )
        payload = repository.get_dashboard_overview(limit_activity=8, limit_expiry=8, critical_days=safe_critical)
        payload["expiry_preview"] = _enrich_verfall_rows(
            payload.get("expiry_preview", []),
            critical_days=safe_critical,
            warning_days=safe_warning,
            attention_days=safe_attention,
        )
        payload["thresholds"] = {
            "kritisch_tage": safe_critical,
            "warnung_tage": safe_warning,
            "achtung_tage": safe_attention,
        }
        return payload

    @app.get("/verfall/overview")
    def verfall_overview(
        perspective: str = "depot",
        ids: str = "",
        q: str = "",
        category: str = "alle",
        limit: int = 100,
        offset: int = 0,
        critical_days: int = 30,
        warning_days: int = 90,
        attention_days: int = 180,
        session: SessionInfo = Depends(require_permission("movements_read")),
    ) -> dict:
        _ = session
        safe_perspective = (perspective or "").strip().lower()
        if safe_perspective not in {"depot", "praeparat"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Perspektive.")
        try:
            selected_ids = _parse_id_list_csv(ids)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Ungueltige IDs: {exc}") from exc
        safe_critical, safe_warning, safe_attention = _normalize_verfall_thresholds(
            critical_days=critical_days,
            warning_days=warning_days,
            attention_days=attention_days,
        )
        depot_ids = selected_ids if safe_perspective == "depot" and selected_ids else None
        praeparat_ids = selected_ids if safe_perspective == "praeparat" and selected_ids else None
        rows = repository.list_verfall_items(
            depot_ids=depot_ids,
            praeparat_ids=praeparat_ids,
            search_text=q,
            category=category,
            limit=limit,
            offset=offset,
            critical_days=safe_critical,
            warning_days=safe_warning,
            attention_days=safe_attention,
        )
        total = repository.count_verfall_items(
            depot_ids=depot_ids,
            praeparat_ids=praeparat_ids,
            search_text=q,
            category=category,
            critical_days=safe_critical,
            warning_days=safe_warning,
            attention_days=safe_attention,
        )
        enriched_rows = _enrich_verfall_rows(rows, safe_critical, safe_warning, safe_attention)
        stats = {
            "kritisch": int(sum(1 for row in enriched_rows if row["kategorie"] == "kritisch")),
            "warnung": int(sum(1 for row in enriched_rows if row["kategorie"] == "warnung")),
            "achtung": int(sum(1 for row in enriched_rows if row["kategorie"] == "achtung")),
            "gesamt_menge": int(sum(int(row.get("anzahl") or 0) for row in enriched_rows)),
        }
        return {
            "rows": enriched_rows,
            "total": total,
            "stats_page": stats,
            "thresholds": {
                "kritisch_tage": safe_critical,
                "warnung_tage": safe_warning,
                "achtung_tage": safe_attention,
            },
        }

    @app.get("/verfall/overview/export.csv")
    def export_verfall_overview_csv(
        perspective: str = "depot",
        ids: str = "",
        q: str = "",
        category: str = "alle",
        critical_days: int = 30,
        warning_days: int = 90,
        attention_days: int = 180,
        session: SessionInfo = Depends(require_permission("movements_read")),
    ) -> StreamingResponse:
        _ = session
        safe_perspective = (perspective or "").strip().lower()
        if safe_perspective not in {"depot", "praeparat"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Perspektive.")
        try:
            selected_ids = _parse_id_list_csv(ids)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Ungueltige IDs: {exc}") from exc
        safe_critical, safe_warning, safe_attention = _normalize_verfall_thresholds(
            critical_days=critical_days,
            warning_days=warning_days,
            attention_days=attention_days,
        )
        depot_ids = selected_ids if safe_perspective == "depot" and selected_ids else None
        praeparat_ids = selected_ids if safe_perspective == "praeparat" and selected_ids else None
        rows = repository.list_verfall_items(
            depot_ids=depot_ids,
            praeparat_ids=praeparat_ids,
            search_text=q,
            category=category,
            limit=5000,
            offset=0,
            critical_days=safe_critical,
            warning_days=safe_warning,
            attention_days=safe_attention,
        )
        enriched_rows = _enrich_verfall_rows(rows, safe_critical, safe_warning, safe_attention)
        csv_rows = [
            [
                row.get("id"),
                row.get("depot"),
                row.get("praeparat"),
                row.get("charge"),
                row.get("verfall"),
                row.get("tage_bis_verfall"),
                row.get("kategorie"),
                row.get("anzahl"),
            ]
            for row in enriched_rows
        ]
        return _csv_response(
            filename="verfall_manager.csv",
            headers=["ID", "Depot", "Praeparat", "Charge", "Verfall", "TageBisVerfall", "Kategorie", "Anzahl"],
            rows=csv_rows,
        )

    @app.get("/notifications/verfall")
    def notifications_verfall(
        since: str = "",
        limit: int = 25,
        critical_days: int = 30,
        warning_days: int = 90,
        attention_days: int = 180,
        session: SessionInfo = Depends(require_permission("movements_read")),
    ) -> dict:
        _ = session
        safe_critical, safe_warning, safe_attention = _normalize_verfall_thresholds(
            critical_days=critical_days,
            warning_days=warning_days,
            attention_days=attention_days,
        )
        now_iso = datetime.now(timezone.utc).isoformat()
        rows = repository.list_new_critical_expiry_events(
            since_iso=(since or "").strip() or None,
            limit=limit,
            critical_days=safe_critical,
        )
        enriched_rows = _enrich_verfall_rows(rows, safe_critical, safe_warning, safe_attention)
        return {
            "since": (since or "").strip() or None,
            "next_since": now_iso,
            "rows": enriched_rows,
            "counts": {
                "kritisch": int(sum(1 for row in enriched_rows if row["kategorie"] == "kritisch")),
                "warnung": int(sum(1 for row in enriched_rows if row["kategorie"] == "warnung")),
                "achtung": int(sum(1 for row in enriched_rows if row["kategorie"] == "achtung")),
                "gesamt": len(enriched_rows),
            },
        }

    @app.post("/bewegungen", status_code=status.HTTP_201_CREATED)
    def create_bewegung(
        payload: BewegungCreateRequest,
        session: SessionInfo = Depends(require_permission("movements_write")),
    ) -> dict[str, int | str]:
        _ = session
        try:
            new_id = repository.insert_bewegung(
                depot_id=payload.depot_id,
                praeparat_id=payload.praeparat_id,
                typ=payload.typ,
                charge=payload.charge,
                verfall=payload.verfall.isoformat(),
                datum=payload.datum.isoformat(),
                anzahl=payload.anzahl,
                empfaenger=payload.empfaenger,
            )
        except ValueError as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=str(exc),
            ) from exc
        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="bewegung",
            resource_id=new_id,
            details={"typ": payload.typ, "depot_id": payload.depot_id, "praeparat_id": payload.praeparat_id},
        )
        return {"id": new_id, "status": "created"}

    @app.post("/bewegungen/{bewegung_id}/attachment", status_code=status.HTTP_201_CREATED)
    def upload_bewegung_attachment(
        bewegung_id: int,
        file: UploadFile = File(...),
        session: SessionInfo = Depends(require_permission("movements_write")),
    ) -> dict[str, int | str]:
        _ = session
        bewegung = repository.get_bewegung(bewegung_id)
        if bewegung is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Bewegung nicht gefunden.")
        filename = (file.filename or "").strip()
        if not filename:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dateiname fehlt.")
        content_type = (file.content_type or "").lower()
        if not filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nur PDF-Dateien sind erlaubt.")
        if content_type and content_type not in PDF_MIME_TYPES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nur PDF-Dateien sind erlaubt.")
        depot_name = repository.get_depot_name(int(bewegung["depot_id"]))
        target_path = _build_attachment_target_path(
            app.state.attachments_dir,
            bewegung_id=bewegung_id,
            depot_name=depot_name,
            original_name=filename,
        )
        try:
            size = _persist_pdf_upload(file, target_path)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        except OSError as exc:
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Datei konnte nicht gespeichert werden.") from exc
        stored_name = Path(filename).name
        uploaded_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        repository.set_bewegung_attachment(
            bewegung_id=bewegung_id,
            datei_pfad=str(target_path),
            datei_name=stored_name,
            datei_groesse=size,
            uploaded_at=uploaded_at,
        )
        repository.log_audit(
            username=session.username,
            action="update",
            resource_type="bewegung",
            resource_id=bewegung_id,
            details={"attachment": stored_name, "size": size},
        )
        return {"id": bewegung_id, "status": "attachment_saved"}

    @app.get("/bewegungen/{bewegung_id}/attachment")
    def get_bewegung_attachment(
        bewegung_id: int,
        download: bool = False,
        session: SessionInfo = Depends(require_permission("movements_read")),
    ) -> FileResponse:
        _ = session
        bewegung = repository.get_bewegung(bewegung_id)
        if bewegung is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Bewegung nicht gefunden.")
        datei_pfad = (bewegung.get("datei_pfad") or "").strip()
        if not datei_pfad:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Kein PDF-Anhang vorhanden.")
        path = Path(datei_pfad)
        if not path.exists() or not path.is_file():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="PDF-Datei nicht gefunden.")
        filename = (bewegung.get("datei_name") or path.name).strip() or path.name
        if download:
            return FileResponse(path=path, media_type="application/pdf", filename=filename)
        return FileResponse(path=path, media_type="application/pdf")

    @app.get("/imports/bewegungen/template")
    def download_bewegungen_template(
        session: SessionInfo = Depends(require_permission("import_use")),
    ) -> StreamingResponse:
        _ = session
        content = _build_import_template(repository)
        return StreamingResponse(
            BytesIO(content),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="bewegungen_vorlage.xlsx"'},
        )

    @app.post("/imports/bewegungen/preview")
    def preview_import_bewegungen(
        file: UploadFile = File(...),
        session: SessionInfo = Depends(require_permission("import_use")),
    ) -> dict:
        _ = session
        filename = (file.filename or "").strip()
        if not filename:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dateiname fehlt.")
        try:
            raw_df, row_offset = _load_import_dataframe(file)
            mapped_df = _map_import_columns(raw_df)
            parsed_rows, preview_rows, errors = _analyze_import_dataframe(repository, mapped_df, row_offset)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Datei konnte nicht gelesen werden: {exc}") from exc
        return {
            "filename": filename,
            "total_rows": int(len(mapped_df.dropna(how="all"))),
            "valid_rows": len(parsed_rows),
            "error_count": len(errors),
            "preview_rows": preview_rows,
            "errors": errors[:100],
        }

    @app.post("/imports/bewegungen/execute")
    def execute_import_bewegungen(
        file: UploadFile = File(...),
        session: SessionInfo = Depends(require_permission("import_use")),
    ) -> dict:
        _ = session
        filename = (file.filename or "").strip()
        if not filename:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dateiname fehlt.")
        try:
            raw_df, row_offset = _load_import_dataframe(file)
            mapped_df = _map_import_columns(raw_df)
            parsed_rows, _preview_rows, errors = _analyze_import_dataframe(repository, mapped_df, row_offset)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Datei konnte nicht gelesen werden: {exc}") from exc

        imported = 0
        for row in parsed_rows:
            try:
                repository.insert_bewegung(
                    depot_id=row["depot_id"],
                    praeparat_id=row["praeparat_id"],
                    typ=row["typ"],
                    charge=row["charge"],
                    verfall=row["verfall"],
                    datum=row["datum"],
                    anzahl=row["anzahl"],
                    empfaenger=row["empfaenger"],
                )
                imported += 1
            except ValueError as exc:
                errors.append(f"Zeile {row['display_row']}: {exc}")

        repository.log_audit(
            username=session.username,
            action="create",
            resource_type="bewegung_import",
            details={"filename": filename, "imported": imported, "errors": len(errors)},
        )
        return {
            "filename": filename,
            "imported": imported,
            "errors": errors[:200],
            "error_count": len(errors),
        }

    @app.get("/reports/bewegungen")
    def report_bewegungen(
        perspective: str = "depot",
        ids: str = "",
        start_date: str = "",
        end_date: str = "",
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> dict:
        _ = session
        safe_perspective = (perspective or "").strip().lower()
        if safe_perspective not in {"depot", "praeparat"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Perspektive.")
        try:
            selected_ids = _parse_id_list_csv(ids)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Ungueltige IDs: {exc}") from exc
        if not selected_ids:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Mindestens eine ID ist erforderlich.")
        rows = repository.get_bewegungen_analyse(
            depot_ids=selected_ids if safe_perspective == "depot" else None,
            praeparat_ids=selected_ids if safe_perspective == "praeparat" else None,
            start_date=(start_date or "").strip() or None,
            end_date=(end_date or "").strip() or None,
        )
        total = int(sum(int(row["anzahl"]) for row in rows))
        zugang = int(sum(int(row["anzahl"]) for row in rows if row["typ"] == "Zugang"))
        abgang = int(sum(int(row["anzahl"]) for row in rows if row["typ"] in {"Abgang", "Vernichtung"}))
        return {
            "rows": rows,
            "kpis": {"gesamt": total, "zugang": zugang, "abgang": abgang},
        }

    @app.get("/reports/bewegungen/export.csv")
    def export_report_bewegungen_csv(
        perspective: str = "depot",
        ids: str = "",
        start_date: str = "",
        end_date: str = "",
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        report = report_bewegungen(perspective, ids, start_date, end_date, session)
        rows = [
            [row["monat"], row["depot"], row["praeparat"], row["typ"], row["anzahl"]]
            for row in report["rows"]
        ]
        return _csv_response(
            filename="bewegungsanalyse.csv",
            headers=["Monat", "Depot", "Praeparat", "Typ", "Anzahl"],
            rows=rows,
        )

    @app.get("/reports/bestand")
    def report_bestand(
        perspective: str = "depot",
        ids: str = "",
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> dict:
        _ = session
        safe_perspective = (perspective or "").strip().lower()
        if safe_perspective not in {"depot", "praeparat"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Perspektive.")
        try:
            selected_ids = _parse_id_list_csv(ids)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Ungueltige IDs: {exc}") from exc
        if not selected_ids:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Mindestens eine ID ist erforderlich.")
        rows = repository.get_bestandsentwicklung(
            depot_ids=selected_ids if safe_perspective == "depot" else None,
            praeparat_ids=selected_ids if safe_perspective == "praeparat" else None,
        )
        sum_soll = int(sum(int(row["sollbestand"]) for row in rows))
        sum_ist = int(sum(int(row["ist_bestand"]) for row in rows))
        luecken = int(sum(1 for row in rows if int(row["differenz"]) < 0))
        quote = (sum_ist / sum_soll * 100) if sum_soll > 0 else 0.0
        return {
            "rows": rows,
            "kpis": {
                "bestandsquote": round(quote, 2),
                "kritische_luecken": luecken,
                "gesamtbestand": sum_ist,
            },
        }

    @app.get("/reports/bestand/export.csv")
    def export_report_bestand_csv(
        perspective: str = "depot",
        ids: str = "",
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        report = report_bestand(perspective, ids, session)
        rows = [
            [row["depot"], row["praeparat"], row["sollbestand"], row["ist_bestand"], row["differenz"]]
            for row in report["rows"]
        ]
        return _csv_response(
            filename="bestandsentwicklung.csv",
            headers=["Depot", "Praeparat", "Soll", "Ist", "Differenz"],
            rows=rows,
        )

    @app.get("/reports/ranking")
    def report_ranking(
        perspective: str = "depot",
        ids: str = "",
        start_date: str = "",
        end_date: str = "",
        limit: int = 10,
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> dict:
        _ = session
        safe_perspective = (perspective or "").strip().lower()
        if safe_perspective not in {"depot", "praeparat"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Perspektive.")
        try:
            selected_ids = _parse_id_list_csv(ids)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Ungueltige IDs: {exc}") from exc
        if not selected_ids:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Mindestens eine ID ist erforderlich.")
        if safe_perspective == "depot":
            rows = repository.get_praeparat_ranking(
                depot_ids=selected_ids,
                start_date=(start_date or "").strip() or None,
                end_date=(end_date or "").strip() or None,
                limit=limit,
            )
            label = "praeparat"
        else:
            rows = repository.get_depot_ranking(
                praeparat_ids=selected_ids,
                start_date=(start_date or "").strip() or None,
                end_date=(end_date or "").strip() or None,
                limit=limit,
            )
            label = "depot"
        total = int(sum(int(row["anzahl"]) for row in rows))
        return {"rows": rows, "label": label, "kpis": {"gesamtvolumen": total}}

    @app.get("/reports/ranking/export.csv")
    def export_report_ranking_csv(
        perspective: str = "depot",
        ids: str = "",
        start_date: str = "",
        end_date: str = "",
        limit: int = 10,
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        report = report_ranking(perspective, ids, start_date, end_date, limit, session)
        rows = [[row["name"], row["anzahl"]] for row in report["rows"]]
        label_header = "Praeparat" if report["label"] == "praeparat" else "Depot"
        return _csv_response(
            filename="ranking.csv",
            headers=[label_header, "Anzahl"],
            rows=rows,
        )

    @app.get("/reports/matrix")
    def report_matrix(
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> dict:
        _ = session
        rows = repository.get_matrix_data()
        enriched = []
        for row in rows:
            diff = int(row["ist_bestand"]) - int(row["sollbestand"])
            item = dict(row)
            item["differenz"] = diff
            enriched.append(item)
        return {"rows": enriched}

    @app.get("/reports/matrix/export.csv")
    def export_report_matrix_csv(
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        report = report_matrix(session)
        rows = [
            [row["depot"], row["praeparat"], row["sollbestand"], row["ist_bestand"], row["differenz"]]
            for row in report["rows"]
        ]
        return _csv_response(
            filename="matrix_bestand.csv",
            headers=["Depot", "Praeparat", "Soll", "Ist", "Differenz"],
            rows=rows,
        )

    @app.get("/reports/verfall")
    def report_verfall(
        perspective: str = "depot",
        ids: str = "",
        horizon_months: int = 24,
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> dict:
        _ = session
        safe_perspective = (perspective or "").strip().lower()
        if safe_perspective not in {"depot", "praeparat"}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ungueltige Perspektive.")
        try:
            selected_ids = _parse_id_list_csv(ids)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Ungueltige IDs: {exc}") from exc
        if not selected_ids:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Mindestens eine ID ist erforderlich.")
        rows = repository.get_verfall_prognose(
            depot_ids=selected_ids if safe_perspective == "depot" else None,
            praeparat_ids=selected_ids if safe_perspective == "praeparat" else None,
            horizon_months=horizon_months,
        )
        current_month = date.today().strftime("%Y-%m")
        horizon_end = _add_months(date.today().replace(day=1), max(1, min(int(horizon_months), 60))).strftime("%Y-%m")
        filtered_rows = [
            row
            for row in rows
            if row["verfall_monat"] and row["verfall_monat"] <= horizon_end
        ]
        bereits_verfallen = int(sum(int(row["anzahl"]) for row in filtered_rows if row["verfall_monat"] < current_month))
        return {
            "rows": filtered_rows,
            "kpis": {
                "bereits_verfallen": bereits_verfallen,
                "gesamt_vorschau": int(sum(int(row["anzahl"]) for row in filtered_rows)),
            },
        }

    @app.get("/reports/verfall/export.csv")
    def export_report_verfall_csv(
        perspective: str = "depot",
        ids: str = "",
        horizon_months: int = 24,
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        report = report_verfall(perspective, ids, horizon_months, session)
        rows = [[row["verfall_monat"], row["anzahl"]] for row in report["rows"]]
        return _csv_response(
            filename="verfall_prognose.csv",
            headers=["Verfall_Monat", "Anzahl"],
            rows=rows,
        )

    @app.get("/reports/{report_type}/export.pdf")
    def export_report_pdf(
        report_type: str,
        perspective: str = "depot",
        ids: str = "",
        start_date: str = "",
        end_date: str = "",
        horizon_months: int = 24,
        limit: int = 10,
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        safe_type = (report_type or "").strip().lower()
        if safe_type == "bewegungen":
            report = report_bewegungen(perspective, ids, start_date, end_date, session)
            headers = ["Monat", "Depot", "Praeparat", "Typ", "Anzahl"]
            rows = [[row["monat"], row["depot"], row["praeparat"], row["typ"], row["anzahl"]] for row in report["rows"]]
            return _pdf_response("bewegungsanalyse.pdf", "Bewegungsanalyse", headers, rows, report["kpis"])
        if safe_type == "bestand":
            report = report_bestand(perspective, ids, session)
            headers = ["Depot", "Praeparat", "Soll", "Ist", "Differenz"]
            rows = [[row["depot"], row["praeparat"], row["sollbestand"], row["ist_bestand"], row["differenz"]] for row in report["rows"]]
            return _pdf_response("bestandsentwicklung.pdf", "Bestandsentwicklung", headers, rows, report["kpis"])
        if safe_type == "ranking":
            report = report_ranking(perspective, ids, start_date, end_date, limit, session)
            label_header = "Praeparat" if report["label"] == "praeparat" else "Depot"
            headers = [label_header, "Anzahl"]
            rows = [[row["name"], row["anzahl"]] for row in report["rows"]]
            return _pdf_response("ranking.pdf", "Ranking", headers, rows, report["kpis"])
        if safe_type == "matrix":
            report = report_matrix(session)
            headers = ["Depot", "Praeparat", "Soll", "Ist", "Differenz"]
            rows = [[row["depot"], row["praeparat"], row["sollbestand"], row["ist_bestand"], row["differenz"]] for row in report["rows"]]
            return _pdf_response("matrix_bestand.pdf", "Matrix Bestand", headers, rows, {})
        if safe_type == "verfall":
            report = report_verfall(perspective, ids, horizon_months, session)
            headers = ["Verfall_Monat", "Anzahl"]
            rows = [[row["verfall_monat"], row["anzahl"]] for row in report["rows"]]
            return _pdf_response("verfall_prognose.pdf", "Verfall Prognose", headers, rows, report["kpis"])
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Unbekannter Report-Typ.")

    @app.get("/reports/{report_type}/export.pptx")
    def export_report_pptx(
        report_type: str,
        perspective: str = "depot",
        ids: str = "",
        start_date: str = "",
        end_date: str = "",
        horizon_months: int = 24,
        limit: int = 10,
        session: SessionInfo = Depends(require_permission("reports_view")),
    ) -> StreamingResponse:
        safe_type = (report_type or "").strip().lower()
        if safe_type == "bewegungen":
            report = report_bewegungen(perspective, ids, start_date, end_date, session)
            headers = ["Monat", "Depot", "Praeparat", "Typ", "Anzahl"]
            rows = [[row["monat"], row["depot"], row["praeparat"], row["typ"], row["anzahl"]] for row in report["rows"]]
            return _pptx_response("bewegungsanalyse.pptx", "Bewegungsanalyse", headers, rows, report["kpis"])
        if safe_type == "bestand":
            report = report_bestand(perspective, ids, session)
            headers = ["Depot", "Praeparat", "Soll", "Ist", "Differenz"]
            rows = [[row["depot"], row["praeparat"], row["sollbestand"], row["ist_bestand"], row["differenz"]] for row in report["rows"]]
            return _pptx_response("bestandsentwicklung.pptx", "Bestandsentwicklung", headers, rows, report["kpis"])
        if safe_type == "ranking":
            report = report_ranking(perspective, ids, start_date, end_date, limit, session)
            label_header = "Praeparat" if report["label"] == "praeparat" else "Depot"
            headers = [label_header, "Anzahl"]
            rows = [[row["name"], row["anzahl"]] for row in report["rows"]]
            return _pptx_response("ranking.pptx", "Ranking", headers, rows, report["kpis"])
        if safe_type == "matrix":
            report = report_matrix(session)
            headers = ["Depot", "Praeparat", "Soll", "Ist", "Differenz"]
            rows = [[row["depot"], row["praeparat"], row["sollbestand"], row["ist_bestand"], row["differenz"]] for row in report["rows"]]
            return _pptx_response("matrix_bestand.pptx", "Matrix Bestand", headers, rows, {})
        if safe_type == "verfall":
            report = report_verfall(perspective, ids, horizon_months, session)
            headers = ["Verfall_Monat", "Anzahl"]
            rows = [[row["verfall_monat"], row["anzahl"]] for row in report["rows"]]
            return _pptx_response("verfall_prognose.pptx", "Verfall Prognose", headers, rows, report["kpis"])
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Unbekannter Report-Typ.")

    @app.get("/audit-logs")
    def list_audit_logs(
        limit: int = 100,
        offset: int = 0,
        q: str = "",
        action: str = "",
        resource_type: str = "",
        session: SessionInfo = Depends(require_permission("audit_view")),
    ) -> list[dict]:
        _ = session
        return repository.list_audit_logs(
            limit=limit,
            offset=offset,
            q=q,
            action=action,
            resource_type=resource_type,
        )

    return app


app = create_app()

